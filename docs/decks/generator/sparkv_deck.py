"""
sparkv_deck.py - Apple-minimal deck generator for SparkV (python-pptx 1.0.2).

One layout model drives three outputs:
  * the .pptx (python-pptx, 13.333 x 7.5 in, blank layouts, native charts),
  * an HTML preview (one 1600 x 900 frame per slide, 120 px per inch),
  * a programmatic layout check (text fit, overlaps, containment, safe area,
    image upscale, minimum type size, red share).

Every element a layout function places is recorded in a registry (El) with its
box in inches, its parent container and the elements it is allowed to sit on
(intentional layers). The checker and the HTML exporter read only that registry,
so the preview and the checks describe exactly what went into the .pptx.

Style source: docs/decks/style-demos/apple-minimal/STYLE.md (tokens in section
STYLE TOKENS below). Units: inches everywhere in the public API, points for type.
"""
from __future__ import annotations

import glob
import html as _html
import json
import math
import os
import re
import shutil
import subprocess
from dataclasses import dataclass, field, replace
from typing import Any, Iterable, Optional, Sequence, Union

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION,
                            XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK)
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

__all__ = ["Deck", "TextFitError", "LayoutError", "accent", "load_outline", "SCREENSHOTS",
           "col_x", "span", "C", "TYPE"]

HERE = os.path.dirname(os.path.abspath(__file__))
DECKS = os.path.dirname(HERE)
SCREENSHOTS = os.path.normpath(os.path.join(DECKS, "..", "screenshots"))

# --------------------------------------------------------------------------- units
EMU_PER_IN = 914400
SLIDE_W_EMU, SLIDE_H_EMU = 12192000, 6858000
SLIDE_W = SLIDE_W_EMU / EMU_PER_IN          # 13.333 in
SLIDE_H = SLIDE_H_EMU / EMU_PER_IN          # 7.5 in
PX_PER_IN = 120                              # HTML frame: 1600 x 900
PT_PER_PX = 0.6


def E(inches: float) -> Emu:
    return Emu(int(round(inches * EMU_PER_IN)))


def _px(inches: float) -> float:
    return round(inches * PX_PER_IN, 2)


# --------------------------------------------------------------------------- STYLE TOKENS
C = {
    # 60: grounds
    "white": "#FFFFFF", "mist": "#F5F5F7",
    # 30: ink and neutrals
    "ink": "#1D1D1F", "ink2": "#6E6E73", "ink3": "#86868B",
    "hairline": "#D2D2D7", "track": "#E8E8ED", "bar": "#AEAEB2",
    # 10: accent
    "red": "#F01822",
    # device
    "bezel": "#0A0A0B", "lid_edge": "#BDBDC2",
    "base_top": "#E3E3E7", "base_mid": "#CFCFD4", "base_bot": "#A8A8AD",
    "thumb_top": "#B4B4B9", "thumb_bot": "#CDCDD2",
    "scr_top": "#2A2E43", "scr_mid": "#171925", "scr_end": "#0F1119",
    "notch": "#000000", "cam_a": "#2B3448", "cam_b": "#11141C",
}


def col(c: Optional[str]) -> Optional[str]:
    if c is None:
        return None
    c = C.get(c, c)
    if not re.fullmatch(r"#[0-9A-Fa-f]{6}", c):
        raise ValueError(f"unknown colour {c!r}")
    return c.upper()


def _rgb(c: str) -> RGBColor:
    return RGBColor.from_string(col(c)[1:])


# Grid (STYLE.md section 4): 12 x 0.7 in columns, 32 px gutters, 1 in margins.
MARGIN = 1.0
COL_W = 0.7
GUTTER = 32 / PX_PER_IN
CONTENT_W = 12 * COL_W + 11 * GUTTER          # 11.333
TOP = 96 / PX_PER_IN                          # 0.8
BOTTOM = 820 / PX_PER_IN                      # 6.833
FOOT_Y = 844 / PX_PER_IN                      # 7.033
FOOT_H = 20 / PX_PER_IN                       # 0.167
SAFE = (MARGIN, 0.75, MARGIN + CONTENT_W, 7.25)   # x0, y0, x1, y1
HAIR_PT = 0.75
EPS = 0.004                                    # in; geometric tolerance


def col_x(c: int) -> float:
    """Left edge (in) of grid column c, 1-based."""
    if not 1 <= c <= 12:
        raise ValueError("column must be 1..12")
    return MARGIN + (c - 1) * (COL_W + GUTTER)


def span(n: int) -> float:
    """Width (in) of n grid columns including the gutters between them."""
    if not 1 <= n <= 12:
        raise ValueError("span must be 1..12")
    return n * COL_W + (n - 1) * GUTTER


@dataclass(frozen=True)
class TS:
    weight: str            # light | regular | medium
    size: float            # pt
    lh: float              # line height as a multiple of size
    tracking: float = 0.0  # em
    color: str = "ink"
    caps: bool = False
    min_scale: float = 0.9


# Type scale (STYLE.md section 3), sizes in pt.
TYPE = {
    "display":   TS("light", 55.2, 1.08, -0.02, "ink", min_scale=0.8),
    "h1":        TS("light", 38.4, 1.12, -0.015, "ink", min_scale=0.8),
    "h1_large":  TS("light", 48.0, 1.16, -0.015, "ink", min_scale=0.8),
    "h1_side":   TS("light", 36.0, 1.12, -0.015, "ink", min_scale=0.8),
    "h1_dash":   TS("light", 33.6, 1.12, -0.015, "ink", min_scale=0.8),
    "eyebrow":   TS("medium", 13.2, 16.8 / 13.2, 0, "ink2"),
    "lead":      TS("regular", 16.8, 1.36, 0, "ink2"),
    "h3":        TS("medium", 16.8, 1.2, 0, "ink"),
    "summary":   TS("medium", 18.0, 1.2, 0, "ink"),
    "body":      TS("regular", 13.2, 1.45, 0, "ink2"),
    "body_ink":  TS("regular", 13.2, 1.45, 0, "ink"),
    "small":     TS("regular", 10.2, 1.45, 0, "ink2"),
    "small_ink": TS("medium", 10.2, 1.45, 0, "ink"),
    "plan_list": TS("regular", 9.6, 1.4, 0, "ink2"),
    "label":     TS("medium", 9.0, 1.2, 0.04, "ink3", caps=True),
    "sub":       TS("regular", 7.8, 1.2, 0, "ink3", min_scale=1.0),
    "footer":    TS("regular", 8.4, 12 / 8.4, 0, "ink3", min_scale=0.93),
    "kpi":       TS("light", 48.0, 1.1, -0.02, "ink", min_scale=0.55),
    "tile":      TS("light", 26.4, 1.12, -0.01, "ink", min_scale=0.7),
    "price":     TS("light", 38.4, 1.12, -0.02, "ink", min_scale=0.7),
    "unit":      TS("regular", 12.0, 1.3, 0, "ink2"),
    "plan_name": TS("medium", 16.8, 1.2, 0, "ink"),
    "tag":       TS("medium", 8.4, 1.2, 0, "red"),
    "inr":       TS("regular", 12.0, 1.3, 0, "ink"),
    "note":      TS("regular", 9.0, 1.4, 0, "ink3"),
    "button":    TS("medium", 12.0, 1.2, 0, "white", min_scale=1.0),
    "cell":      TS("regular", 10.2, 1.35, 0, "ink2"),
    "cell_ink":  TS("medium", 10.2, 1.35, 0, "ink"),
    "cell_sm":   TS("regular", 9.0, 1.35, 0, "ink2", min_scale=0.87),
    "cell_sm_ink": TS("medium", 9.0, 1.35, 0, "ink", min_scale=0.87),
    "quote":     TS("light", 33.6, 1.2, -0.01, "ink", min_scale=0.75),
}
MIN_FONT_PT = 7.8

# --------------------------------------------------------------------------- fonts
_HN = "/System/Library/Fonts/HelveticaNeue.ttc"


def _sf_file(name: str) -> Optional[str]:
    for d in ("/Library/Fonts", os.path.expanduser("~/Library/Fonts")):
        p = os.path.join(d, name)
        if os.path.exists(p):
            return p
    return None


def _profiles() -> dict:
    profs = {
        # Helvetica Neue ships with macOS; measurement uses the same faces.
        "mac": {
            "light": ("Helvetica Neue Light", (_HN, 7), (_HN, 8), 300),
            "regular": ("Helvetica Neue", (_HN, 0), (_HN, 2), 400),
            "medium": ("Helvetica Neue Medium", (_HN, 10), (_HN, 11), 500),
            "css": '"Helvetica Neue", Helvetica, Arial, sans-serif',
            "css_link": "", "safety": 1.04,
        },
        # Windows-safe deck: Inter everywhere (measured with Helvetica Neue plus a
        # wider safety margin, Inter runs about 5% wider).
        "inter": {
            "light": ("Inter Light", (_HN, 7), (_HN, 8), 300),
            "regular": ("Inter", (_HN, 0), (_HN, 2), 400),
            "medium": ("Inter SemiBold", (_HN, 10), (_HN, 11), 600),
            "css": '"Inter", "Helvetica Neue", Arial, sans-serif',
            "css_link": '<link href="https://fonts.googleapis.com/css2?family=Inter:ital,wght@0,300;0,400;0,600;1,300;1,400;1,600&display=swap" rel="stylesheet">',
            "safety": 1.10,
        },
    }
    lt, rg, sb = (_sf_file("SF-Pro-Display-Light.otf"), _sf_file("SF-Pro-Text-Regular.otf"),
                  _sf_file("SF-Pro-Display-Semibold.otf"))
    if lt and rg and sb:
        lti = _sf_file("SF-Pro-Display-LightItalic.otf") or lt
        rgi = _sf_file("SF-Pro-Text-RegularItalic.otf") or rg
        sbi = _sf_file("SF-Pro-Display-SemiboldItalic.otf") or sb
        profs["sf"] = {
            "light": ("SF Pro Display Light", (lt, 0), (lti, 0), 300),
            "regular": ("SF Pro Text", (rg, 0), (rgi, 0), 400),
            "medium": ("SF Pro Display Semibold", (sb, 0), (sbi, 0), 600),
            "css": '-apple-system, "SF Pro Display", "SF Pro Text", "Helvetica Neue", sans-serif',
            "css_link": "", "safety": 1.04,
        }
    return profs


PROFILES = _profiles()
_FALLBACK_FONTS = ["/System/Library/Fonts/SFNS.ttf", "/System/Library/Fonts/Apple Symbols.ttf",
                   "/Library/Fonts/Arial Unicode.ttf"]
_CMAP_CACHE: dict = {}


def _cmap(path: str, index: int) -> set:
    key = (path, index)
    if key not in _CMAP_CACHE:
        from fontTools.ttLib import TTCollection, TTFont
        if path.lower().endswith(".ttc"):
            f = TTCollection(path).fonts[index]
        else:
            f = TTFont(path, fontNumber=index)
        _CMAP_CACHE[key] = set((f.getBestCmap() or {}).keys())
    return _CMAP_CACHE[key]


class Metrics:
    """Text measurement with PIL ImageFont, per weight/italic face, with glyph fallback."""

    def __init__(self, profile: dict):
        self.p = profile
        self.safety = profile["safety"]
        self._faces: dict = {}
        self._cache: dict = {}
        self._fallbacks = []
        for p in _FALLBACK_FONTS:
            if os.path.exists(p):
                try:
                    self._fallbacks.append((ImageFont.truetype(p, 100), _cmap(p, 0)))
                except Exception:
                    pass

    def _face(self, weight: str, italic: bool):
        key = (weight, italic)
        if key not in self._faces:
            _, reg, ita, _ = self.p[weight]
            path, idx = ita if italic else reg
            self._faces[key] = (ImageFont.truetype(path, 100, index=idx), _cmap(path, idx))
        return self._faces[key]

    def _w100(self, text: str, weight: str, italic: bool) -> float:
        key = (text, weight, italic)
        if key in self._cache:
            return self._cache[key]
        font, cmap = self._face(weight, italic)
        total, buf = 0.0, ""
        for ch in text:
            if ord(ch) in cmap or ch.isspace():
                buf += ch
                continue
            if buf:
                total += font.getlength(buf)
                buf = ""
            fb = next((f for f, cm in self._fallbacks if ord(ch) in cm), None)
            total += (fb or font).getlength(ch) * (1.0 if fb else 1.2)
        if buf:
            total += font.getlength(buf)
        self._cache[key] = total
        return total

    def width(self, text: str, weight: str, italic: bool, size: float, tracking: float = 0.0) -> float:
        """Width in pt, safety margin included."""
        if not text:
            return 0.0
        return self._w100(text, weight, italic) * size / 100.0 * self.safety + tracking * size * len(text)


# --------------------------------------------------------------------------- text model
@dataclass
class Run:
    text: str
    weight: str = "regular"
    size: float = 13.2
    color: str = "#6E6E73"
    italic: bool = False
    tracking: float = 0.0


@dataclass
class Para:
    runs: list
    align: str = "l"           # l | c | r
    lh: float = 1.45
    space_before: float = 0.0  # pt

    @property
    def size(self) -> float:
        return max((r.size for r in self.runs), default=10.0)

    @property
    def pitch(self) -> float:
        return self.size * self.lh

    def scaled(self, k: float) -> "Para":
        return Para([replace(r, size=round(r.size * k, 2)) for r in self.runs], self.align, self.lh,
                    round(self.space_before * k, 2))

    def plain(self) -> str:
        return "".join(r.text for r in self.runs)


_MARK = re.compile(r"(\*\*.+?\*\*|==.+?==|\[TBD[^\]]*\])")
TBD_RE = re.compile(r"\[TBD[^\]]*\]")


def accent(header: str, header_accent: Optional[str]) -> str:
    """Wrap the accent phrase of a header in **...** (semibold emphasis in Apple minimal)."""
    if not header_accent or header_accent not in header:
        return header
    i = header.index(header_accent)
    j = i + len(header_accent)
    # keep trailing punctuation inside the emphasis, like the style demo
    while j < len(header) and header[j] in ".,!?:;":
        j += 1
    return header[:i] + "**" + header[i:j] + "**" + header[j:]


class TextFitError(Exception):
    pass


class LayoutError(Exception):
    pass


# --------------------------------------------------------------------------- registry
@dataclass
class El:
    eid: int
    slide: int
    kind: str                      # text | rect | line | image | chart | mark | group
    x: float
    y: float
    w: float
    h: float
    parent: Optional[int] = None   # container element id; None = slide
    hosts: tuple = ()              # elements this one is intentionally layered on
    bleed: bool = False            # may leave the safe area (device bleeding off the slide)
    name: str = ""
    props: dict = field(default_factory=dict)

    @property
    def r(self) -> float:
        return self.x + self.w

    @property
    def b(self) -> float:
        return self.y + self.h

    def box(self) -> tuple:
        return (round(self.x, 4), round(self.y, 4), round(self.w, 4), round(self.h, 4))


@dataclass
class SlideRec:
    index: int
    layout: str
    bg: str
    page: str
    hidden: bool
    notes: str
    slide: Any = None
    els: list = field(default_factory=list)


# --------------------------------------------------------------------------- helpers
def _nice_max(v: float) -> float:
    if v <= 0:
        return 1.0
    v *= 1.12
    e = 10 ** math.floor(math.log10(v))
    for m in (1, 1.2, 1.5, 2, 2.5, 3, 4, 5, 6, 8, 10):
        if m * e >= v:
            return m * e
    return 10 * e


def fmt_num(v: Optional[float], f: str) -> str:
    """Format a number the way an Excel/PowerPoint number format would ($#,##0.00, 0.0"M", 0%)."""
    if v is None:
        return ""
    if f in ("", "General"):
        return f"{v:g}"
    m = re.match(r'^(?P<pre>[^#0]*)(?P<num>[#,0]*0(?:\.0+)?)(?P<suf>.*)$', f)
    if not m:
        return f"{v:g}"
    pre, num, suf = m["pre"].replace('"', ""), m["num"], m["suf"].replace('"', "")
    if "%" in suf:
        v = v * 100
    dec = len(num.split(".")[1]) if "." in num else 0
    s = f"{abs(v):,.{dec}f}" if "," in num else f"{abs(v):.{dec}f}"
    return ("-" if v < 0 else "") + pre + s + suf


def load_outline(deck: str) -> dict:
    with open(os.path.join(DECKS, deck, "outline.json"), encoding="utf-8") as f:
        return json.load(f)


def _remove_style(shape) -> None:
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def _set_grad(spPr, stops: Sequence, angle_deg: float = 90.0) -> None:
    """Replace the fill in spPr with a linear gradient. angle 90 = top to bottom."""
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    g = etree.Element(qn("a:gradFill"))
    g.set("rotWithShape", "1")
    lst = etree.SubElement(g, qn("a:gsLst"))
    for pos, c in stops:
        gs = etree.SubElement(lst, qn("a:gs"))
        gs.set("pos", str(int(round(max(0.0, min(1.0, pos)) * 100000))))
        clr = etree.SubElement(gs, qn("a:srgbClr"))
        clr.set("val", col(c)[1:])
    lin = etree.SubElement(g, qn("a:lin"))
    lin.set("ang", str(int(angle_deg * 60000)))
    lin.set("scaled", "0")
    # fill must come right after the geometry
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    geom.addnext(g)


# =========================================================================== Deck
class Deck:
    """A SparkV deck in the Apple-minimal style.

    Deck(title, font_profile='auto'|'mac'|'inter'|'sf', ref_ppi=144, max_upscale=1.25)

    ref_ppi: pixel density at which a screenshot counts as shown at 1.0x. 144 px/in is a
    13.333 in slide played at 1920 x 1080. max_upscale caps the displayed size of any
    image at max_upscale x its native pixels at that density.
    """

    def __init__(self, title: str = "SparkV", *, font_profile: str = "auto",
                 screenshots_dir: str = SCREENSHOTS, ref_ppi: float = 144.0,
                 max_upscale: float = 1.25, footer_default: str = "", author: str = "SparkV"):
        if font_profile == "auto":
            font_profile = "sf" if "sf" in PROFILES else "mac"
        if font_profile not in PROFILES:
            raise ValueError(f"font profile {font_profile!r} not available: {sorted(PROFILES)}")
        self.profile_name = font_profile
        self.fp = PROFILES[font_profile]
        self.m = Metrics(self.fp)
        self.screens = screenshots_dir
        self.ref_ppi = ref_ppi
        self.max_upscale = max_upscale
        self.footer_default = footer_default
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W_EMU
        self.prs.slide_height = SLIDE_H_EMU
        self.prs.core_properties.title = title
        self.prs.core_properties.author = author
        self._blank = self.prs.slide_layouts[6]
        self.title = title
        self.slides: list = []
        self._eid = 0
        self._img_cache: dict = {}

    # ------------------------------------------------------------------ registry
    def _reg(self, s: SlideRec, kind: str, x, y, w, h, *, parent=None, hosts=(), bleed=False,
             name="", props=None) -> El:
        self._eid += 1
        if w < -EPS or h < -EPS:
            raise LayoutError(f"slide {s.page}: negative size for {name or kind}")
        hp = tuple(hosts) if isinstance(hosts, (list, tuple)) else (hosts,)
        el = El(self._eid, s.index, kind, float(x), float(y), float(max(w, 0)), float(max(h, 0)),
                parent.eid if isinstance(parent, El) else parent,
                tuple(h_.eid if isinstance(h_, El) else h_ for h_ in hp if h_ is not None),
                bleed, name, props or {})
        s.els.append(el)
        return el

    # ------------------------------------------------------------------ text building
    def P(self, text: Union[str, Para], style: str = "body", *, align: str = "l", color: Optional[str] = None,
          weight: Optional[str] = None, space_before: float = 0.0, size: Optional[float] = None) -> Para:
        """Build a paragraph from a string with inline markup.

        **text**  emphasis (semibold family, ink)   ==text==  SparkV red
        [TBD: ...]  placeholder, always grey italic (ink3)  \\n  line break
        """
        if isinstance(text, Para):
            return text
        ts = TYPE[style]
        base_color = col(color or ts.color)
        size = size or ts.size
        runs = []
        for part in _MARK.split(text):
            if not part:
                continue
            w, c, it, t = weight or ts.weight, base_color, False, part
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                t, w = part[2:-2], "medium"
                if base_color in (col("ink2"), col("ink3")):
                    c = col("ink")
            elif part.startswith("==") and part.endswith("==") and len(part) > 4:
                t, c = part[2:-2], col("red")
            elif TBD_RE.fullmatch(part):
                c, it = col("ink3"), True
            if ts.caps:
                t = t.upper()
            runs.append(Run(t, w, size, c, it, ts.tracking))
        if not runs:
            runs = [Run("", weight or ts.weight, size, base_color, False, ts.tracking)]
        return Para(runs, align, ts.lh, space_before)

    def _paras(self, content, style, align, color, space) -> list:
        items = content if isinstance(content, (list, tuple)) else [content]
        out = []
        for i, it in enumerate(items):
            p = self.P(it, style, align=align, color=color) if not isinstance(it, Para) else it
            if i > 0 and space and not p.space_before:
                p = Para(p.runs, p.align, p.lh, space)
            out.append(p)
        return out

    def wrap(self, para: Para, width_pt: float):
        """Greedy word wrap. Returns (line_count, overflow_word or None, lines as text)."""
        tokens = []
        for r in para.runs:
            for part in re.split(r"(\n| +)", r.text):
                if part == "":
                    continue
                if part == "\n":
                    tokens.append(("br", None))
                elif part.isspace():
                    tokens.append(("s", [(part, r)]))
                elif tokens and tokens[-1][0] == "w":
                    tokens[-1][1].append((part, r))
                else:
                    tokens.append(("w", [(part, r)]))
        lines, cur, pend, overflow = 1, 0.0, 0.0, None
        texts, cur_t, pend_t = [], "", ""
        for kind, segs in tokens:
            if kind == "br":
                texts.append(cur_t)
                lines, cur, pend, cur_t, pend_t = lines + 1, 0.0, 0.0, "", ""
                continue
            wdt = sum(self.m.width(t, r.weight, r.italic, r.size, r.tracking) for t, r in segs)
            txt = "".join(t for t, _ in segs)
            if kind == "s":
                if cur > 0:
                    pend += wdt
                    pend_t += txt
                continue
            if cur == 0:
                cur, cur_t = wdt, txt
            elif cur + pend + wdt <= width_pt:
                cur += pend + wdt
                cur_t += pend_t + txt
            else:
                texts.append(cur_t)
                lines += 1
                cur, cur_t = wdt, txt
            if wdt > width_pt + 0.01:
                overflow = txt
            pend, pend_t = 0.0, ""
        texts.append(cur_t)
        return lines, overflow, texts

    def measure(self, paras: Sequence[Para], width_in: float):
        """Return (height_in, overflow_word, line_counts) for paragraphs in a box width_in wide."""
        wpt = width_in * 72.0
        h, ov, counts = 0.0, None, []
        for i, p in enumerate(paras):
            n, o, _ = self.wrap(p, wpt)
            counts.append(n)
            ov = ov or o
            h += (p.space_before if i else 0.0) + n * p.pitch
        return h / 72.0, ov, counts

    def fit_text(self, paras: Sequence[Para], w: float, h: float, min_scale: float = 0.9,
                 where: str = "") -> tuple:
        """Shrink fonts in steps (down to min_scale) until paras fit in w x h inches.
        Returns (scaled paras, needed height in, scale). Raises TextFitError otherwise."""
        k = 1.0
        while True:
            ps = [p.scaled(k) for p in paras] if k != 1.0 else list(paras)
            need, ov, _ = self.measure(ps, w)
            if need <= h + 1e-6 and not ov and min(r.size for p in ps for r in p.runs) >= MIN_FONT_PT - 1e-6:
                return ps, need, k
            nk = round(k - 0.025, 3)
            if nk < min_scale - 1e-9:
                txt = " / ".join(p.plain() for p in paras)
                why = f"word {ov!r} wider than {w:.3f} in" if ov else f"needs {need:.3f} in, box {h:.3f} in"
                raise TextFitError(f"{where}: text does not fit at {k:.0%} scale ({why}): {txt[:140]!r}")
            k = nk

    # ------------------------------------------------------------------ primitives
    def text(self, s: SlideRec, x: float, y: float, w: float, h: Optional[float], content,
             style: str = "body", *, align: str = "l", anchor: str = "t", parent=None, hosts=(),
             max_h: Optional[float] = None, min_scale: Optional[float] = None, name: str = "",
             color: Optional[str] = None, space: float = 0.0, fill: Optional[str] = None,
             radius: float = 0.0, inset: tuple = (0, 0, 0, 0), bleed: bool = False) -> El:
        """Place a text box. h=None sizes the box to its measured content (limited by max_h).
        Fonts shrink within min_scale if needed; TextFitError if the text still cannot fit."""
        paras = self._paras(content, style, align, color, space)
        ms = min_scale if min_scale is not None else TYPE[style].min_scale
        il, it, ir, ib = inset
        iw = w - il - ir
        where = f"slide {s.page} {name or style}"
        if h is None:
            limit = (max_h - it - ib) if max_h is not None else 1e9
            paras, need, k = self.fit_text(paras, iw, limit, ms, where)
            h = need + it + ib + 0.0005
        else:
            paras, need, k = self.fit_text(paras, iw, h - it - ib, ms, where)
        if fill is not None:
            shp = s.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                           E(x), E(y), E(w), E(h))
            _remove_style(shp)
            shp.fill.solid()
            shp.fill.fore_color.rgb = _rgb(fill)
            shp.line.fill.background()
            if radius:
                shp.adjustments[0] = min(0.5, radius / min(w, h))
            tf = shp.text_frame
        else:
            shp = s.slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
            tf = shp.text_frame
        shp.name = (name or style)[:60]
        self._write_tf(tf, paras, anchor, inset)
        return self._reg(s, "text", x, y, w, h, parent=parent, hosts=hosts, bleed=bleed, name=name or style,
                         props={"paras": paras, "anchor": anchor, "fill": col(fill) if fill else None,
                                "radius": radius, "inset": inset, "need": need + it + ib, "scale": k})

    def _write_tf(self, tf, paras, anchor, inset):
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        il, it, ir, ib = inset
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = E(il), E(it), E(ir), E(ib)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[p.align]
            para.line_spacing = Pt(round(p.pitch, 2))
            para.space_before = Pt(p.space_before if i else 0)
            para.space_after = Pt(0)
            for r in p.runs:
                for j, piece in enumerate(r.text.split("\n")):
                    if j:
                        para.add_line_break()
                    if not piece:
                        continue
                    run = para.add_run()
                    run.text = piece
                    f = run.font
                    f.name = self.fp[r.weight][0]
                    f.size = Pt(r.size)
                    f.bold = False
                    f.italic = bool(r.italic)
                    f.color.rgb = _rgb(r.color)
                    if r.tracking:
                        run._r.get_or_add_rPr().set("spc", str(int(round(r.tracking * r.size * 100))))
            # end-of-paragraph size, so an empty line keeps the right height
            endp = para._p.get_or_add_endParaRPr()
            endp.set("sz", str(int(round(p.size * 100))))

    def rect(self, s: SlideRec, x, y, w, h, *, fill: Optional[str] = None, grad: Optional[Sequence] = None,
             line: Optional[str] = None, line_w: float = HAIR_PT, radius: float = 0.0,
             radii: Optional[tuple] = None, oval: bool = False, parent=None, hosts=(), bleed=False,
             name: str = "") -> El:
        """Rectangle / rounded rectangle / oval. radii=(top, bottom) rounds the two pairs separately."""
        if oval:
            st = MSO_SHAPE.OVAL
        elif radii is not None:
            st = MSO_SHAPE.ROUND_2_SAME_RECTANGLE
        elif radius:
            st = MSO_SHAPE.ROUNDED_RECTANGLE
        else:
            st = MSO_SHAPE.RECTANGLE
        shp = s.slide.shapes.add_shape(st, E(x), E(y), E(w), E(h))
        _remove_style(shp)
        shp.name = (name or "rect")[:60]
        m = min(w, h) or 1
        if radii is not None:
            shp.adjustments[0] = min(0.5, radii[0] / m)
            shp.adjustments[1] = min(0.5, radii[1] / m)
        elif radius:
            shp.adjustments[0] = min(0.5, radius / m)
        if grad:
            _set_grad(shp._element.spPr, grad, 90)
        elif fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = _rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = _rgb(line)
            shp.line.width = Pt(line_w)
        else:
            shp.line.fill.background()
        if radii is not None:
            rr = (radii[0], radii[0], radii[1], radii[1])
        else:
            rr = (radius,) * 4
        return self._reg(s, "rect", x, y, w, h, parent=parent, hosts=hosts, bleed=bleed, name=name or "rect",
                         props={"fill": col(fill) if fill else None,
                                "grad": [(p, col(c)) for p, c in grad] if grad else None,
                                "line": col(line) if line else None, "line_w": line_w,
                                "radii": rr, "oval": oval})

    def line(self, s: SlideRec, x1, y1, x2, y2, *, color: str = "hairline", width: float = HAIR_PT,
             parent=None, hosts=(), name: str = "rule") -> El:
        c = s.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        _remove_style(c)
        c.line.color.rgb = _rgb(color)
        c.line.width = Pt(width)
        c.name = name[:60]
        return self._reg(s, "line", min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1), parent=parent,
                         hosts=hosts, name=name, props={"color": col(color), "width": width,
                                                        "pts": (x1, y1, x2, y2)})

    def mark(self, s: SlideRec, kind: str, x, y, size: float = 0.14, *, color: str = "ink",
             width: float = 1.5, parent=None, name: str = "") -> El:
        """Check (kind='check') or cross (kind='cross') mark drawn with two strokes in a group."""
        if kind == "check":
            pts = [((0.05, 0.55), (0.38, 0.86)), ((0.38, 0.86), (0.95, 0.16))]
        elif kind == "cross":
            pts = [((0.12, 0.12), (0.88, 0.88)), ((0.88, 0.12), (0.12, 0.88))]
        else:
            raise ValueError("mark kind must be check or cross")
        grp = s.slide.shapes.add_group_shape()
        grp.name = f"{kind} mark"
        for (ax, ay), (bx, by) in pts:
            c = grp.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x + ax * size), E(y + ay * size),
                                         E(x + bx * size), E(y + by * size))
            _remove_style(c)
            c.line.color.rgb = _rgb(color)
            c.line.width = Pt(width)
        return self._reg(s, "mark", x, y, size, size, parent=parent, name=name or kind,
                         props={"mark": kind, "color": col(color), "width": width, "segs": pts})

    # ------------------------------------------------------------------ images
    def image_path(self, name: str, final: Optional[str] = None, media_root: Optional[str] = None) -> str:
        """Resolve a screenshot name; prefer the recorded 'final' still if it exists."""
        if final and media_root:
            fp = os.path.join(media_root, final)
            if os.path.exists(fp):
                return fp
        p = name if os.path.isabs(name) else os.path.join(self.screens, name)
        if not os.path.exists(p):
            raise FileNotFoundError(f"image not found: {p}")
        return p

    def image_size(self, path: str) -> tuple:
        if path not in self._img_cache:
            with Image.open(path) as im:
                self._img_cache[path] = im.size
        return self._img_cache[path]

    def edge_colors(self, path: str) -> tuple:
        """Average colour of the screenshot's top rows and bottom rows (hex), for a seamless wallpaper."""
        with Image.open(path) as im:
            im = im.convert("RGB")
            w, h = im.size
            def avg(box):
                px = list(im.crop(box).resize((1, 1), Image.Resampling.BOX).getdata())[0]
                return "#%02X%02X%02X" % px
            return avg((0, 0, w, max(1, h // 60))), avg((0, h - max(1, h // 60), w, h))

    def max_width_in(self, path: str) -> float:
        """Largest display width (in) that keeps the image within max_upscale at ref_ppi."""
        return self.image_size(path)[0] * self.max_upscale / self.ref_ppi

    def fit_image(self, path: str, x, y, w, h, *, mode: str = "contain", align: str = "c",
                  valign: str = "m") -> dict:
        """Place an image in a box keeping its aspect ratio.
        contain: whole image inside the box. cover: fill the box, cropping the overflow.
        Never shows the image larger than max_upscale (it is centred smaller instead).
        Returns dict(x, y, w, h, crop=(l, t, r, b))."""
        iw, ih = self.image_size(path)
        ar = iw / ih
        maxw = self.max_width_in(path)
        crop = (0.0, 0.0, 0.0, 0.0)
        if mode == "contain":
            dw = min(w, h * ar, maxw)
            dh = dw / ar
        elif mode == "cover":
            if w / h > ar:      # box wider: crop top/bottom
                dw = min(w, maxw)
                dh = min(h, dw / ar)
                vis = dh / (dw / ar)
                crop = (0.0, 0.0, 0.0, 1 - vis) if valign == "t" else (0.0, (1 - vis) / 2, 0.0, (1 - vis) / 2)
            else:               # box taller: crop left/right
                dh = h
                dw_full = dh * ar
                if dw_full > maxw:
                    dw_full = maxw
                    dh = dw_full / ar
                dw = min(w, dw_full)
                vis = dw / dw_full
                crop = ((1 - vis) / 2, 0.0, (1 - vis) / 2, 0.0)
        else:
            raise ValueError("mode must be contain or cover")
        dx = {"l": 0, "c": (w - dw) / 2, "r": w - dw}[align]
        dy = {"t": 0, "m": (h - dh) / 2, "b": h - dh}[valign]
        return {"x": x + dx, "y": y + dy, "w": dw, "h": dh, "crop": crop}

    def picture(self, s: SlideRec, path: str, x, y, w, h, *, crop=(0, 0, 0, 0), radius: float = 0.0,
                parent=None, hosts=(), bleed=False, name: str = "") -> El:
        pic = s.slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
        pic.name = (name or os.path.basename(path))[:60]
        l, t, r, b = crop
        if any(crop):
            pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b
        if radius:
            geom = pic._element.spPr.find(qn("a:prstGeom"))
            geom.set("prst", "roundRect")
            av = geom.find(qn("a:avLst"))
            if av is None:
                av = etree.SubElement(geom, qn("a:avLst"))
            for gd in list(av):
                av.remove(gd)
            gd = etree.SubElement(av, qn("a:gd"))
            gd.set("name", "adj")
            gd.set("fmla", f"val {int(min(0.5, radius / min(w, h)) * 100000)}")
        iw, ih = self.image_size(path)
        return self._reg(s, "image", x, y, w, h, parent=parent, hosts=hosts, bleed=bleed,
                         name=name or os.path.basename(path),
                         props={"src": path, "crop": tuple(crop), "radius": radius, "px": (iw, ih)})

    # ------------------------------------------------------------------ device mockup
    def macbook(self, s: SlideRec, x: float, y: float, lid_w: float, image: str, *,
                pad_px: float = 12, shot_frac: float = 1.0, clip_bottom: Optional[float] = None,
                with_base: bool = True, name: str = "device", wallpaper: str = "auto") -> dict:
        """MacBook notch mockup drawn with shapes (STYLE.md section 5), back to front:
        lid, screen (wallpaper gradient; wallpaper='auto' samples the screenshot's own top and
        bottom edge colours so the shot blends with no seam, 'style' uses the STYLE.md stops), screenshot (top flush, centred), notch with
        camera dot, base slab with thumb cut. clip_bottom cuts the device at that y
        (bleeding off the slide edge); the base is then omitted.
        Returns the element boxes."""
        path = self.image_path(image)
        iw, ih = self.image_size(path)
        edge = 2 / PX_PER_IN
        pad = pad_px / PX_PER_IN
        pe = pad + edge
        scr_w = lid_w - 2 * pe
        scr_h = scr_w / 1.54
        lid_h = pe + scr_h + pad + 6 / PX_PER_IN
        bleed = False
        if clip_bottom is not None and y + lid_h > clip_bottom:
            lid_h = clip_bottom - y
            bleed = True
            with_base = False
        vis_scr_h = min(scr_h, y + lid_h - (y + pe))
        lid = self.rect(s, x, y, lid_w, lid_h, fill="bezel", line="lid_edge", line_w=1.2,
                        radii=(26 / PX_PER_IN, 0 if bleed else 6 / PX_PER_IN), bleed=bleed,
                        name=f"{name} lid")
        # screenshot geometry first (the screen gradient stops at its bottom edge)
        shot_w = min(scr_w * shot_frac, self.max_width_in(path))
        shot_h = shot_w * ih / iw
        if shot_h > vis_scr_h:
            shot_h = vis_scr_h
            shot_w = shot_h * iw / ih
        sx = x + pe
        sy = y + pe
        stop = shot_h / vis_scr_h
        top_c, bot_c = self.edge_colors(path) if wallpaper == "auto" else (C["scr_top"], C["scr_mid"])
        screen = self.rect(s, sx, sy, scr_w, vis_scr_h,
                           grad=[(0, top_c), (stop, bot_c), (1, "scr_end")],
                           radii=(10 / PX_PER_IN, 0 if bleed else 2 / PX_PER_IN), parent=lid, bleed=bleed,
                           name=f"{name} screen")
        shot = self.picture(s, path, sx + (scr_w - shot_w) / 2, sy, shot_w, shot_h, parent=screen,
                            bleed=bleed, name=f"{name} screenshot {os.path.basename(path)}")
        nw, nh = 0.30 * shot_w, 0.0444 * shot_w
        notch = self.rect(s, sx + (scr_w - nw) / 2, sy, nw, nh, fill="notch", radii=(0, 14 / PX_PER_IN * nw / 2.083),
                          parent=screen, hosts=(shot,), bleed=bleed, name=f"{name} notch")
        dot = min(10 / PX_PER_IN, nh * 0.45)
        self.rect(s, sx + scr_w / 2 - dot / 2, sy + nh / 2 - dot / 2, dot, dot, oval=True,
                  grad=[(0, "cam_a"), (1, "cam_b")], parent=notch, bleed=bleed, name=f"{name} camera")
        out = {"lid": lid, "screen": screen, "shot": shot, "notch": notch, "bottom": y + lid_h}
        if with_base:
            bw = lid_w * (7.467 / 6.833)
            bh = 22 / PX_PER_IN
            bx = x + (lid_w - bw) / 2
            base = self.rect(s, bx, y + lid_h, bw, bh, grad=[(0, "base_top"), (0.45, "base_mid"), (1, "base_bot")],
                             radii=(3 / PX_PER_IN, 20 / PX_PER_IN), name=f"{name} base")
            tw = 0.16 * bw
            self.rect(s, bx + (bw - tw) / 2, y + lid_h, tw, 9 / PX_PER_IN, grad=[(0, "thumb_top"), (1, "thumb_bot")],
                      radii=(0, 10 / PX_PER_IN), parent=base, name=f"{name} thumb cut")
            out["base"] = base
            out["bottom"] = y + lid_h + bh
        return out

    @staticmethod
    def macbook_height(lid_w: float, pad_px: float = 12, with_base: bool = True) -> float:
        pe = pad_px / PX_PER_IN + 2 / PX_PER_IN
        scr_h = (lid_w - 2 * pe) / 1.54
        return pe + scr_h + pad_px / PX_PER_IN + 6 / PX_PER_IN + (22 / PX_PER_IN if with_base else 0)

    @staticmethod
    def macbook_lid_for_height(total_h: float, pad_px: float = 12) -> float:
        pe = pad_px / PX_PER_IN + 2 / PX_PER_IN
        fixed = pe + pad_px / PX_PER_IN + 6 / PX_PER_IN + 22 / PX_PER_IN
        return (total_h - fixed) * 1.54 + 2 * pe

    def window_shot(self, s: SlideRec, image: str, x, y, w, h, *, radius: float = 0.1,
                    name: str = "window") -> El:
        """A macOS window screenshot, contained in the box with rounded corners, no border."""
        path = self.image_path(image)
        f = self.fit_image(path, x, y, w, h, mode="contain")
        return self.picture(s, path, f["x"], f["y"], f["w"], f["h"], radius=radius, name=f"{name} {image}")

    # ------------------------------------------------------------------ slide scaffolding
    def _new(self, layout: str, bg: str, notes: str, footer, page: Optional[str], hidden: bool) -> SlideRec:
        slide = self.prs.slides.add_slide(self._blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(bg)
        idx = len(self.slides) + 1
        rec = SlideRec(idx, layout, col(bg), page or f"{idx:02d}", hidden, notes or "", slide)
        self.slides.append(rec)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        if hidden:
            slide._element.set("show", "0")
        if footer is not False:
            self._footer(rec, footer if footer is not None else self.footer_default)
        return rec

    def _footer(self, s: SlideRec, text: str) -> None:
        num_w = 0.6
        if text:
            self.text(s, MARGIN, FOOT_Y, CONTENT_W - num_w - GUTTER, FOOT_H, text, "footer", name="footer")
        self.text(s, MARGIN + CONTENT_W - num_w, FOOT_Y, num_w, FOOT_H, s.page, "footer", align="r",
                  name="page number")

    def _header(self, s: SlideRec, eyebrow: Optional[str], title: Optional[str], lead: Optional[str] = None,
                *, x: float = MARGIN, w: float = CONTENT_W, align: str = "l", h1: str = "h1",
                title_lines: int = 2, lead_lines: int = 2, eyebrow_color: Optional[str] = None,
                top: float = TOP, gap_eyebrow: float = 16 / PX_PER_IN, gap_lead: float = 26 / PX_PER_IN) -> float:
        """Eyebrow, headline and lead stacked from `top`. Returns the bottom y."""
        y = top
        if eyebrow:
            e = self.text(s, x, y, w, None, eyebrow, "eyebrow", align=align, color=eyebrow_color,
                          max_h=TYPE["eyebrow"].size * TYPE["eyebrow"].lh / 72 + 0.001, name="eyebrow")
            y = e.b + gap_eyebrow
        if title:
            ts = TYPE[h1]
            t = self.text(s, x, y, w, None, title, h1, align=align,
                          max_h=title_lines * ts.size * ts.lh / 72 + 0.001, name="headline")
            y = t.b
        if lead:
            ts = TYPE["lead"]
            l_ = self.text(s, x, y + gap_lead, w, None, lead, "lead", align=align,
                           max_h=lead_lines * ts.size * ts.lh / 72 + 0.001, name="lead")
            y = l_.b
        return y

    def button(self, s: SlideRec, x: float, y: float, text: str, *, parent=None, name: str = "button") -> El:
        """Red pill button (padding 15 x 30 px). Returns the element; width follows the text."""
        ts = TYPE["button"]
        p = self.P(text, "button", align="c")
        tw = sum(self.m.width(r.text, r.weight, r.italic, r.size, r.tracking) for r in p.runs) / 72
        padx, pady = 30 / PX_PER_IN, 15 / PX_PER_IN
        w = tw + 2 * padx + 0.02
        h = ts.size * ts.lh / 72 + 2 * pady
        return self.text(s, x, y, w, h, [p], "button", anchor="m", fill="red", radius=h / 2,
                         inset=(padx, pady, padx, pady), parent=parent, name=name)

    def button_width(self, text: str) -> float:
        p = self.P(text, "button")
        return sum(self.m.width(r.text, r.weight, r.italic, r.size, r.tracking) for r in p.runs) / 72 \
            + 2 * 30 / PX_PER_IN + 0.02

    def text_width(self, text: str, style: str) -> float:
        """Single-line width (in) of marked-up text in a style."""
        p = self.P(text, style)
        return sum(self.m.width(r.text, r.weight, r.italic, r.size, r.tracking) for r in p.runs) / 72

    def card(self, s: SlideRec, x, y, w, h, *, fill: str = "mist", radius: float = 22 / PX_PER_IN,
             line: Optional[str] = None, line_w: float = HAIR_PT, name: str = "card") -> El:
        return self.rect(s, x, y, w, h, fill=fill, radius=radius, line=line, line_w=line_w, name=name)

    # ================================================================== layouts
    def title_slide(self, title: str, *, lead: Optional[str] = None, eyebrow: Optional[str] = "SparkV",
                    image: Optional[str] = "01-home.png", cta: Optional[str] = None, notes: str = "",
                    lid_w: float = 9.167, page: Optional[str] = None, shot_frac: float = 1.0) -> SlideRec:
        """Slide 1: centred red eyebrow, display headline (2 lines), lead, optional pill,
        MacBook bleeding off the bottom edge. No footer."""
        s = self._new("title_slide", "mist", notes, False, page, False)
        y = TOP
        if eyebrow:
            e = self.text(s, MARGIN, y, CONTENT_W, None, eyebrow, "eyebrow", align="c", color="red", name="eyebrow")
            y = e.b + 18 / PX_PER_IN
        ts = TYPE["display"]
        t = self.text(s, MARGIN, y, CONTENT_W, None, title, "display", align="c",
                      max_h=2 * ts.size * ts.lh / 72 + 0.001, name="headline")
        y = t.b
        if lead:
            l_ = self.text(s, MARGIN, y + 26 / PX_PER_IN, CONTENT_W, None, lead, "lead", align="c",
                           max_h=2 * TYPE["lead"].size * TYPE["lead"].lh / 72 + 0.001, name="lead")
            y = l_.b
        if cta:
            bw = self.button_width(cta)
            b = self.button(s, (SLIDE_W - bw) / 2, y + 0.25, cta, name="cta")
            y = b.b
        if image:
            dy = max(3.767, y + 0.35)
            if dy + 1.0 > SLIDE_H:
                raise LayoutError("title_slide: not enough room for the device")
            self.macbook(s, (SLIDE_W - lid_w) / 2, dy, lid_w, image, pad_px=16, shot_frac=shot_frac,
                         clip_bottom=SLIDE_H, name="hero device")
        return s

    def section_header(self, title: str, *, eyebrow: Optional[str] = None, lead: Optional[str] = None,
                       image: Optional[str] = None, captions: Sequence[str] = (), notes: str = "",
                       footer=None, page: Optional[str] = None, hidden: bool = False) -> SlideRec:
        """Key-visual header: centred eyebrow/headline/lead, MacBook in the middle,
        up to 3 captions on 4-column spans along the bottom, each under a hairline."""
        s = self._new("section_header", "mist", notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, align="c", h1="h1", title_lines=2)
        cap_top = BOTTOM
        if captions:
            if len(captions) > 3:
                raise LayoutError("section_header: at most 3 captions")
            n = len(captions)
            spans = {1: [(1, 12)], 2: [(1, 6), (7, 6)], 3: [(1, 4), (5, 4), (9, 4)]}[n]
            paras = [self._paras(c, "body", "l", None, 0) for c in captions]
            need = max(self.measure(p, span(w_))[0] for p, (_, w_) in zip(paras, spans))
            cap_top = BOTTOM - need - 0.15
            for c, (c0, w_) in zip(captions, spans):
                self.line(s, col_x(c0), cap_top, col_x(c0) + span(w_), cap_top, name="caption rule")
                self.text(s, col_x(c0), cap_top + 0.15, span(w_), need, c, "body", name="caption")
        if image:
            avail_top, avail_bot = y + 0.3, cap_top - 0.3
            lid = min(self.macbook_lid_for_height(avail_bot - avail_top), 6.833)
            if lid < 3.0:
                raise LayoutError(f"section_header slide {s.page}: no room for the device ({lid:.2f} in lid)")
            h = self.macbook_height(lid)
            dy = avail_top + (avail_bot - avail_top - h) / 2
            self.macbook(s, (SLIDE_W - lid) / 2, dy, lid, image, shot_frac=1.0, name="device")
        return s

    def minimal_body(self, title: str, *, eyebrow: Optional[str] = None, lead: Optional[str] = None,
                     points: Sequence = (), numbered: bool = True, notes: str = "", footer=None,
                     page: Optional[str] = None, hidden: bool = False, bg: str = "white") -> SlideRec:
        """Problem-style slide: large headline on 10 columns, then 2-4 points bottom-aligned
        on equal spans, each under a hairline: number label, optional H3, text.
        points: strings, or (h3, body) tuples."""
        s = self._new("minimal_body", bg, notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, w=span(10), h1="h1_large", title_lines=2)
        n = len(points)
        if not 1 <= n <= 4:
            raise LayoutError("minimal_body: 1 to 4 points")
        spans = {1: [(1, 8)], 2: [(1, 6), (7, 6)], 3: [(1, 4), (5, 4), (9, 4)],
                 4: [(1, 3), (4, 3), (7, 3), (10, 3)]}[n]
        blocks = []
        for i, (pt_, (c0, w_)) in enumerate(zip(points, spans)):
            h3, body = (pt_ if isinstance(pt_, (tuple, list)) else (None, pt_))
            items = []
            if numbered:
                items.append(("label", f"{i + 1:02d}", 0))
            if h3:
                items.append(("h3", h3, 18 / PX_PER_IN))
            items.append(("body" if h3 else "lead", body, (14 if h3 else 18) / PX_PER_IN))
            hs = [self.measure([self.P(t, st)], span(w_))[0] for st, t, _ in items]
            total = 28 / PX_PER_IN + sum(hs) + sum(g for _, _, g in items[1:])
            blocks.append((items, hs, total, c0, w_))
        tallest = max(b[2] for b in blocks)
        top = BOTTOM - tallest
        if top < y + 0.4:
            raise LayoutError(f"minimal_body slide {s.page}: points collide with the headline")
        for items, hs, total, c0, w_ in blocks:
            x, w = col_x(c0), span(w_)
            self.line(s, x, top, x + w, top, name="point rule")
            yy = top + 28 / PX_PER_IN
            for k, ((st, t, gap), hh) in enumerate(zip(items, hs)):
                if k:
                    yy += gap
                self.text(s, x, yy, w, hh + 0.0005, t, st, name=f"point {st}")
                yy += hh + 0.0005
        return s

    def feature_showcase(self, title: str, *, eyebrow: Optional[str] = None, lead: Optional[str] = None,
                         bullets: Sequence = (), image: Optional[str] = None, frame: str = "notch",
                         notes: str = "", footer=None, page: Optional[str] = None,
                         hidden: bool = False, shot_frac: float = 1.0) -> SlideRec:
        """Text on columns 1-4 (vertically centred), screenshot on columns 5-12.
        frame='notch' (or 'band'): MacBook notch mockup. frame='window': rounded window shot.
        bullets: strings, or (key, label) tuples for a hotkey list (key column 0.967 in)."""
        s = self._new("feature_showcase", "mist", notes, footer, page, hidden)
        x, w = col_x(1), span(4)
        # measure the text block first so it can be centred
        parts = []
        if eyebrow:
            parts.append(("eyebrow", eyebrow, 0))
        parts.append(("h1_side", title, 20 / PX_PER_IN if eyebrow else 0))
        if lead:
            parts.append(("body", lead, 24 / PX_PER_IN))
        row_pad = 14 / PX_PER_IN
        key_w = 0.967

        def row_h(b, k=1.0):
            if isinstance(b, (tuple, list)):
                hk = self.measure([self.P(b[0], "small_ink").scaled(k)], key_w - 0.05)[0]
                hl = self.measure([self.P(b[1], "small").scaled(k)], w - key_w - 16 / PX_PER_IN)[0]
                return max(hk, hl) + 2 * row_pad
            return self.measure([self.P(b, "small").scaled(k)], w)[0] + 2 * row_pad

        def block_h(k=1.0):
            hh = 0.0
            for st, t, gap in parts:
                lim = 3 if st == "h1_side" else 4
                ts = TYPE[st]
                h_ = min(self.measure([self.P(t, st)], w)[0], lim * ts.size * ts.lh / 72)
                hh += gap + h_
            if bullets:
                hh += 36 / PX_PER_IN + sum(row_h(b, k) for b in bullets)
            return hh

        total = block_h()
        avail = BOTTOM - TOP
        if total > avail:
            raise LayoutError(f"feature_showcase slide {s.page}: text block {total:.2f} in > {avail:.2f} in")
        y = TOP + (avail - total) / 2
        for st, t, gap in parts:
            y += gap
            ts = TYPE[st]
            lim = 3 if st == "h1_side" else 4
            el = self.text(s, x, y, w, None, t, st, max_h=lim * ts.size * ts.lh / 72 + 0.001, name=st)
            y = el.b
        if bullets:
            y += 36 / PX_PER_IN
            self.line(s, x, y, x + w, y, name="list rule")
            for b in bullets:
                rh = row_h(b)
                if isinstance(b, (tuple, list)):
                    self.text(s, x, y + row_pad, key_w - 0.05, rh - 2 * row_pad, b[0], "small_ink", name="key")
                    self.text(s, x + key_w + 16 / PX_PER_IN, y + row_pad, w - key_w - 16 / PX_PER_IN,
                              rh - 2 * row_pad, b[1], "small", name="key label")
                else:
                    self.text(s, x, y + row_pad, w, rh - 2 * row_pad, b, "small", name="bullet")
                y += rh
                self.line(s, x, y, x + w, y, name="list rule")
        if image:
            dx, dw = col_x(5), span(8)
            if frame in ("notch", "band", "drawer"):
                lid = 6.833
                h = self.macbook_height(lid)
                self.macbook(s, dx + (dw - lid) / 2, TOP + (avail - h) / 2, lid, image, shot_frac=shot_frac,
                             name="device")
            elif frame == "window":
                self.window_shot(s, image, dx, TOP + 0.1, dw, avail - 0.2)
            else:
                raise LayoutError(f"unknown frame {frame!r}")
        return s

    def _metric_tiles(self, s, metrics, boxes, value_style="kpi"):
        """Mist cards with a big number and a caption; one common number size for all tiles."""
        padx, pady = 32 / PX_PER_IN, 28 / PX_PER_IN
        if min(bx[3] for bx in boxes) < 1.6:
            pady = 20 / PX_PER_IN
        gap = 14 / PX_PER_IN
        ts = TYPE[value_style]
        k = 1.0
        for m, (x, y, w, h) in zip(metrics, boxes):
            p = self._metric_value_para(m, value_style)
            _, _, kk = self.fit_text([p], w - 2 * padx, ts.size * ts.lh / 72 + 0.001, ts.min_scale,
                                     f"slide {s.page} metric {m.get('value')}")
            k = min(k, kk)
            # the number also shrinks (down to min_scale) to leave room for the caption
            lab = self.measure([self.P(m.get("label") or "", "small")], w - 2 * padx)[0] if m.get("label") else 0
            room = h - 2 * pady - (gap + lab if lab else 0)
            k = min(k, room / (ts.size * ts.lh / 72))
        if k < ts.min_scale - 1e-9:
            raise TextFitError(f"slide {s.page}: metric tiles too small for their numbers and captions "
                               f"(scale {k:.2f} < {ts.min_scale})")
        for m, (x, y, w, h) in zip(metrics, boxes):
            cardel = self.card(s, x, y, w, h, name="metric card")
            p = self._metric_value_para(m, value_style).scaled(k)
            vh = p.pitch / 72 + 0.0005
            self.text(s, x + padx, y + pady, w - 2 * padx, vh, [p], value_style, parent=cardel, name="metric value")
            cap_y = y + pady + vh + gap
            if m.get("label"):
                self.text(s, x + padx, cap_y, w - 2 * padx, y + h - pady - cap_y, m["label"], "small",
                          parent=cardel, name="metric label")

    def _metric_value_para(self, m: dict, style: str) -> Para:
        v = str(m.get("value", ""))
        if m.get("accent"):
            return self.P(v, style, color="red", weight="medium")
        return self.P(v, style)

    def metrics_dashboard(self, title: str, *, eyebrow: Optional[str] = None, lead: Optional[str] = None,
                          metrics: Sequence[dict] = (), chart: Optional[dict] = None, notes: str = "",
                          footer=None, page: Optional[str] = None, hidden: bool = False) -> SlideRec:
        """Big-number tiles. Without a chart: 2-6 tiles in a grid (3 across for 3, 5, 6).
        With a chart: 1-3 tiles stacked on columns 1-4, chart card on columns 5-12.
        metric: {value, label, accent: bool}. '[TBD...]' values render grey italic."""
        s = self._new("metrics_dashboard", "white", notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, h1="h1_dash", title_lines=2, lead_lines=1)
        top = y + 0.3
        n = len(metrics)
        if chart:
            if not 0 <= n <= 3:
                raise LayoutError("metrics_dashboard with chart: 0 to 3 metrics")
            gap = 24 / PX_PER_IN
            if n:
                ch = (BOTTOM - top - (n - 1) * gap) / n
                boxes = [(col_x(1), top + i * (ch + gap), span(4), ch) for i in range(n)]
                self._metric_tiles(s, metrics, boxes)
                cx, cw = col_x(5), span(8)
            else:
                cx, cw = col_x(1), span(12)
            self.chart_card(s, cx, top, cw, BOTTOM - top, chart)
        else:
            if not 2 <= n <= 6:
                raise LayoutError("metrics_dashboard: 2 to 6 metrics")
            cols = 3 if n in (3, 5, 6) else 2
            rows = math.ceil(n / cols)
            gap = 24 / PX_PER_IN
            cw = span(12 // cols)
            chh = min(2.3, (BOTTOM - top - (rows - 1) * gap) / rows)
            boxes = []
            for i in range(n):
                r, c = divmod(i, cols)
                boxes.append((col_x(1 + c * (12 // cols)), top + r * (chh + gap), cw, chh))
            self._metric_tiles(s, metrics, boxes)
        return s

    # ------------------------------------------------------------------ charts
    def _norm_chart(self, spec: dict) -> dict:
        sp = dict(spec)
        t = sp.get("type", "column")
        if t not in ("column", "column_stacked", "bar", "bar_stacked", "line", "doughnut"):
            raise ValueError(f"chart type {t!r} not supported")
        sp["type"] = t
        cats = list(sp["categories"])
        series = []
        palette = ["bar", "ink", "hairline", "ink3"]
        for i, ser in enumerate(sp["series"]):
            vals = list(ser["values"])
            if len(vals) != len(cats):
                raise ValueError(f"series {ser.get('name')} has {len(vals)} values for {len(cats)} categories")
            series.append({"name": ser.get("name", f"Series {i + 1}"), "values": vals,
                           "color": col(ser.get("color", palette[i % len(palette)])),
                           "point_colors": {int(k): col(v) for k, v in (ser.get("point_colors") or {}).items()}})
        sp["series"] = series
        sp["categories"] = cats
        sp.setdefault("number_format", "General")
        if t == "doughnut":
            if len(series) != 1:
                raise ValueError("doughnut takes one series")
            dpal = ["red", "ink", "bar", "hairline", "ink3", "track"]
            pc = series[0]["point_colors"]
            for i in range(len(cats)):
                pc.setdefault(i, col(dpal[i % len(dpal)]))
        else:
            vals = [v for ser in series for v in ser["values"] if v is not None]
            if t.endswith("stacked"):
                tot = [sum((ser["values"][i] or 0) for ser in series) for i in range(len(cats))]
                hi, lo = max(tot + [0]), min(tot + [0])
            else:
                hi, lo = max(vals + [0]), min(vals + [0])
            sp["axis_max"] = sp.get("axis_max") or _nice_max(hi)
            sp["axis_min"] = sp.get("axis_min") if sp.get("axis_min") is not None else (0 if lo >= 0 else -_nice_max(-lo))
        sp["legend"] = sp.get("legend", len(series) > 1 and t != "doughnut")
        sp["_norm"] = True
        return sp

    def chart_card(self, s: SlideRec, x, y, w, h, spec: dict, *, name: str = "chart card") -> El:
        """Mist card: title row, native chart, [TBD] notes, source line under the chart."""
        sp = self._norm_chart(spec)
        cardel = self.card(s, x, y, w, h, name=name)
        padx, pady = 32 / PX_PER_IN, 28 / PX_PER_IN
        ix, iw = x + padx, w - 2 * padx
        yy = y + pady
        if sp.get("title"):
            sub = sp.get("subtitle")
            sub_w = self.text_width(sub, "small") + 0.05 if sub else 0
            if sub:
                self.text(s, ix + iw - sub_w, yy, sub_w, None, sub, "small", align="r", color="ink3",
                          parent=cardel, name="chart subtitle")
            t = self.text(s, ix, yy, iw - sub_w - (0.2 if sub else 0), None, sp["title"], "small_ink",
                          max_h=2 * 10.2 * 1.45 / 72 + 0.001, parent=cardel, name="chart title")
            yy = t.b + 0.12
        bottom = y + h - pady
        foot = []
        for tb in sp.get("tbd") or []:
            foot.append(self.P(tb if TBD_RE.search(tb) else f"[TBD: {tb}]", "sub"))
        if sp.get("source"):
            foot.append(self.P("Source: " + sp["source"] if not sp["source"].lower().startswith("source") else sp["source"], "sub"))
        if foot:
            fh = self.measure(foot, iw)[0]
            ftop = bottom - fh
            self.text(s, ix, ftop - 0.0005, iw, fh + 0.0005, foot, "sub", parent=cardel, name="chart source")
            bottom = ftop - 0.1
        ch_h = bottom - yy
        if ch_h < 1.2:
            raise LayoutError(f"slide {s.page}: chart area only {ch_h:.2f} in tall")
        self.chart(s, ix, yy, iw, ch_h, sp, parent=cardel)
        return cardel

    def chart(self, s: SlideRec, x, y, w, h, spec: dict, *, parent=None) -> El:
        """Native python-pptx chart with data labels. Doughnut gets a text legend on the right."""
        sp = spec if spec.get("_norm") else self._norm_chart(spec)
        t = sp["type"]
        font_name = self.fp["regular"][0]
        if t == "doughnut":
            d = min(h, w * 0.45)
            el = self._native_chart(s, x, y + (h - d) / 2, d, d, sp, font_name, parent)
            # legend rows (category, value) to the right
            lx = x + d + 0.35
            lw = min(x + w - lx, 4.0)
            ser = sp["series"][0]
            rows = [(c, ser["values"][i], ser["point_colors"][i]) for i, c in enumerate(sp["categories"])]
            rh = 0.36
            ly = y + (h - rh * len(rows)) / 2
            for i, (c, v, clr) in enumerate(rows):
                ry = ly + i * rh
                self.rect(s, lx, ry + 0.06, 0.12, 0.12, fill=clr, radius=0.02, parent=parent, name="legend swatch")
                self.text(s, lx + 0.22, ry, lw - 0.22 - 1.1, rh - 0.02, str(c), "small", parent=parent,
                          name="legend label")
                self.text(s, lx + lw - 1.0, ry, 1.0, rh - 0.02, fmt_num(v, sp["number_format"]) or "[TBD]",
                          "small_ink", align="r", parent=parent, name="legend value")
            return el
        return self._native_chart(s, x, y, w, h, sp, font_name, parent)

    def _native_chart(self, s, x, y, w, h, sp, font_name, parent) -> El:
        t = sp["type"]
        xl = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
              "bar": XL_CHART_TYPE.BAR_CLUSTERED, "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
              "line": XL_CHART_TYPE.LINE_MARKERS, "doughnut": XL_CHART_TYPE.DOUGHNUT}[t]
        cd = CategoryChartData(number_format=sp["number_format"])
        cd.categories = sp["categories"]
        for ser in sp["series"]:
            cd.add_series(ser["name"], ser["values"])
        gf = s.slide.shapes.add_chart(xl, E(x), E(y), E(w), E(h), cd)
        gf.name = "chart"
        ch = gf.chart
        ch.has_title = False
        ch.font.name = font_name
        ch.font.size = Pt(9)
        ch.font.color.rgb = _rgb("ink2")
        cs = ch._chartSpace
        # transparent chart and plot area (the card shows through), square corners
        sppr = etree.Element(qn("c:spPr"))
        etree.SubElement(sppr, qn("a:noFill"))
        ln = etree.SubElement(sppr, qn("a:ln"))
        etree.SubElement(ln, qn("a:noFill"))
        cs.find(qn("c:chart")).addnext(sppr)
        pa = cs.find(qn("c:chart")).find(qn("c:plotArea"))
        psp = etree.SubElement(pa, qn("c:spPr"))
        etree.SubElement(psp, qn("a:noFill"))
        plot = ch.plots[0]
        if t == "doughnut":
            ch.has_legend = False
            plot.vary_by_categories = True
            ser = plot.series[0]
            for i, c in sp["series"][0]["point_colors"].items():
                pt = ser.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = _rgb(c)
                pt.format.line.color.rgb = _rgb("white")
                pt.format.line.width = Pt(1.5)
            hs = cs.find(".//" + qn("c:holeSize"))
            if hs is None:
                dn = cs.find(".//" + qn("c:doughnutChart"))
                hs = etree.SubElement(dn, qn("c:holeSize"))
            hs.set("val", "62")
        else:
            plot.vary_by_categories = False
            if t in ("column", "bar"):
                plot.gap_width = 70
                plot.overlap = -12 if len(sp["series"]) > 1 else 0
            elif t.endswith("stacked"):
                plot.gap_width = 70
                plot.overlap = 100
            for i, ser in enumerate(plot.series):
                spec_s = sp["series"][i]
                if t == "line":
                    ser.smooth = False
                    ser.format.line.color.rgb = _rgb(spec_s["color"])
                    ser.format.line.width = Pt(2.25)
                    ser.marker.style = XL_MARKER_STYLE.CIRCLE
                    ser.marker.size = 7
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = _rgb(spec_s["color"])
                    ser.marker.format.line.color.rgb = _rgb(spec_s["color"])
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = _rgb(spec_s["color"])
                    ser.format.line.fill.background()
                    ser.invert_if_negative = False
                for pi, pc in spec_s["point_colors"].items():
                    pt = ser.points[pi]
                    if t == "line":
                        pt.marker.format.fill.solid()
                        pt.marker.format.fill.fore_color.rgb = _rgb(pc)
                    else:
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = _rgb(pc)
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = sp["number_format"]
            dl.number_format_is_linked = False
            dl.show_value = True
            dl.font.size = Pt(9)
            dl.font.name = self.fp["medium"][0]
            dl.font.color.rgb = _rgb("ink")
            dl.position = {"column": XL_LABEL_POSITION.OUTSIDE_END, "bar": XL_LABEL_POSITION.OUTSIDE_END,
                           "line": XL_LABEL_POSITION.ABOVE}.get(t, XL_LABEL_POSITION.CENTER)
            va = ch.value_axis
            va.minimum_scale = sp["axis_min"]
            va.maximum_scale = sp["axis_max"]
            va.has_major_gridlines = False
            va.has_minor_gridlines = False
            va.visible = False
            ca = ch.category_axis
            ca.major_tick_mark = XL_TICK_MARK.NONE
            ca.minor_tick_mark = XL_TICK_MARK.NONE
            ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
            ca.has_major_gridlines = False
            ca.format.line.color.rgb = _rgb("hairline")
            ca.format.line.width = Pt(HAIR_PT)
            ca.tick_labels.font.size = Pt(9)
            ca.tick_labels.font.color.rgb = _rgb("ink2")
            if t.startswith("bar"):
                ca.reverse_order = True
            ch.has_legend = bool(sp["legend"])
            if sp["legend"]:
                ch.legend.position = XL_LEGEND_POSITION.BOTTOM
                ch.legend.include_in_layout = False
                ch.legend.font.size = Pt(9)
                ch.legend.font.color.rgb = _rgb("ink2")
        return self._reg(s, "chart", x, y, w, h, parent=parent, name=f"{t} chart", props={"spec": sp})

    def chart_slide(self, title: str, chart: dict, *, eyebrow: Optional[str] = None,
                    lead: Optional[str] = None, side: Sequence[str] = (), notes: str = "", footer=None,
                    page: Optional[str] = None, hidden: bool = False) -> SlideRec:
        """Headline plus one chart card. side: up to 4 short notes on columns 9-12
        (the chart then takes columns 1-8)."""
        s = self._new("chart_slide", "white", notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        top = y + 0.35
        if side:
            self.chart_card(s, col_x(1), top, span(8), BOTTOM - top, chart)
            x, w = col_x(9), span(4)
            yy = top
            for i, t in enumerate(side):
                self.line(s, x, yy, x + w, yy, name="side rule")
                el = self.text(s, x, yy + 18 / PX_PER_IN, w, None, t, "body",
                               max_h=BOTTOM - yy - 18 / PX_PER_IN, name="side note")
                yy = el.b + 24 / PX_PER_IN
            if yy - 24 / PX_PER_IN > BOTTOM + EPS:
                raise LayoutError(f"chart_slide {s.page}: side notes overflow")
        else:
            self.chart_card(s, col_x(1), top, span(12), BOTTOM - top, chart)
        return s

    # ------------------------------------------------------------------ tables
    def _table(self, s: SlideRec, x: float, y: float, w: float, columns: Sequence[str], rows: Sequence[Sequence],
               *, widths: Optional[Sequence[float]] = None, highlight_row: Optional[int] = None,
               max_bottom: float = BOTTOM, cell: str = "cell", first: str = "cell_ink",
               head_us: Optional[int] = None, min_row_h: float = 0.36, inset: float = 0.12,
               vpad: float = 0.09) -> float:
        """Grid table from text boxes and hairlines. Cell values: str, True (check), False (cross),
        None (empty) or (bool, str). Fonts shrink together if the table is too tall.
        Returns the bottom y."""
        n = len(columns)
        if any(len(r) != n for r in rows):
            raise LayoutError("table: every row needs one value per column")
        widths = list(widths or [1] * n)
        cgap = 0.2
        avail = w - 2 * inset - cgap * (n - 1)
        tot = float(sum(widths))
        cw = [avail * f / tot for f in widths]
        cx = [x + inset + sum(cw[:i]) + cgap * i for i in range(n)]
        mark_sz = 0.13
        head_h = TYPE["label"].size * TYPE["label"].lh / 72 + 0.0005

        def cell_paras(v, j, ri, k):
            st = first if j == 0 else cell
            if ri == highlight_row:
                st = first
            if isinstance(v, tuple):
                v = v[1]
            if v is None or isinstance(v, bool):
                return None
            return [self.P(str(v), st).scaled(k)]

        def text_w(v, j):
            return cw[j] - ((mark_sz + 0.08) if (isinstance(v, bool) or isinstance(v, tuple)) else 0)

        def heights(k):
            hs = []
            for ri, r in enumerate(rows):
                hh = 0.0
                for j, v in enumerate(r):
                    ps = cell_paras(v, j, ri, k)
                    if ps:
                        hh = max(hh, self.measure(ps, text_w(v, j))[0])
                hs.append(max(min_row_h, hh + 2 * vpad))
            return hs

        k = 1.0
        ms = min(TYPE[cell].min_scale, TYPE[first].min_scale)
        while True:
            hs = heights(k)
            if y + head_h + 0.12 + sum(hs) <= max_bottom + 1e-6:
                break
            k = round(k - 0.025, 3)
            if k < ms - 1e-9:
                raise TextFitError(f"slide {s.page}: table needs {head_h + 0.12 + sum(heights(ms)):.2f} in, "
                                   f"has {max_bottom - y:.2f} in")
        for j, c in enumerate(columns):
            if c:
                self.text(s, cx[j], y, cw[j], head_h, c, "label", color="ink" if head_us == j else None,
                          name="column head")
        yy = y + head_h + 0.12
        self.line(s, x, yy, x + w, yy, color="ink", name="table head rule")
        for ri, (r, rh) in enumerate(zip(rows, hs)):
            parent = None
            if ri == highlight_row:
                parent = self.rect(s, x, yy, w, rh, fill="mist", name="highlight row")
                self.rect(s, x, yy, 0.05, rh, fill="red", parent=parent, name="highlight accent")
            for j, v in enumerate(r):
                tx = cx[j]
                mk = v[0] if isinstance(v, tuple) else v if isinstance(v, bool) else None
                if mk is not None:
                    self.mark(s, "check" if mk else "cross", tx, yy + vpad + 0.02, mark_sz,
                              color="ink" if mk else "ink3", parent=parent, name="cell mark")
                    tx += mark_sz + 0.08
                ps = cell_paras(v, j, ri, k)
                if ps:
                    self.text(s, tx, yy + vpad, text_w(v, j), rh - 2 * vpad, ps, cell, parent=parent,
                              name=f"cell r{ri + 1}c{j + 1}")
            yy += rh
            if ri != len(rows) - 1:
                self.line(s, x, yy, x + w, yy, name="row rule")
        return yy

    def two_column_comparison(self, title: str, *, left_head: str, right_head: str, rows: Sequence,
                              eyebrow: Optional[str] = None, lead: Optional[str] = None,
                              summary: Optional[tuple] = None, notes: str = "", footer=None,
                              page: Optional[str] = None, hidden: bool = False) -> SlideRec:
        """Label column (cols 1-4) + two option columns (cols 5-8, 9-12).
        rows: (label, left, right); values are str, True (check), False (cross) or (bool, str).
        summary: (label, left text, right text); the right text is SparkV red at 18 pt."""
        s = self._new("two_column_comparison", "mist", notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        top = y + 0.35
        sum_h = 0.0
        sum_paras = None
        if summary:
            sum_paras = [[self.P(summary[0], "h3")], [self.P(summary[1], "body_ink")],
                         [self.P(summary[2], "summary", color="red")]]
            sum_h = max(self.measure(p, span(4) - 0.12)[0] for p in sum_paras) + 0.25
        bottom = self._table(s, MARGIN, top, CONTENT_W, ["", left_head, right_head], rows,
                             widths=[4, 4, 4], max_bottom=BOTTOM - sum_h, head_us=2, inset=0.0,
                             cell="small", first="small_ink", min_row_h=0.34, vpad=0.075)
        if summary:
            self.line(s, MARGIN, bottom, MARGIN + CONTENT_W, bottom, color="ink", name="summary rule")
            cw = (CONTENT_W - 0.4) / 3
            for j, ps in enumerate(sum_paras):
                self.text(s, MARGIN + j * (cw + 0.2), bottom + 0.15, cw, None, ps, "body", name="summary")
        return s

    def comparison_table(self, title: str, *, columns: Sequence[str], rows: Sequence[Sequence],
                         eyebrow: Optional[str] = None, lead: Optional[str] = None,
                         widths: Optional[Sequence[float]] = None, highlight_row: Optional[int] = None,
                         notes: str = "", footer=None, page: Optional[str] = None,
                         hidden: bool = False) -> SlideRec:
        """Competitor table across 12 columns. highlight_row (0-based) gets a mist band and a red edge."""
        s = self._new("comparison_table", "white", notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        self._table(s, MARGIN, y + 0.35, CONTENT_W, columns, rows, widths=widths, highlight_row=highlight_row)
        return s

    def appendix_table(self, title: str, *, columns: Sequence[str], rows: Sequence[Sequence],
                       eyebrow: Optional[str] = "Appendix", lead: Optional[str] = None,
                       widths: Optional[Sequence[float]] = None, highlight_row: Optional[int] = None,
                       notes: str = "", footer=None, page: Optional[str] = None,
                       hidden: bool = False) -> SlideRec:
        """Dense reference table (9 pt cells) for backup slides."""
        s = self._new("appendix_table", "white", notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, h1="h1_dash", title_lines=1, lead_lines=1)
        self._table(s, MARGIN, y + 0.3, CONTENT_W, columns, rows, widths=widths, highlight_row=highlight_row,
                    cell="cell_sm", first="cell_sm_ink", min_row_h=0.3, inset=0.1)
        return s

    # ------------------------------------------------------------------ pricing
    def pricing_table(self, title: str, plans: Sequence[dict], *, eyebrow: Optional[str] = None,
                      lead: Optional[str] = None, cta: Optional[str] = None, aside: Optional[str] = None,
                      notes: str = "", footer=None, page: Optional[str] = None,
                      hidden: bool = False) -> SlideRec:
        """2-4 plan cards. plan: {name, price, unit, inr, inr_aside, note, features: [..],
        highlight: bool, tag}. The highlighted plan gets a 2 px red outline and a red tag.
        cta: red pill text; aside: grey text beside it (either may be None)."""
        s = self._new("pricing_table", "mist", notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, align="c", h1="h1", title_lines=1, lead_lines=1)
        n = len(plans)
        if not 2 <= n <= 4:
            raise LayoutError("pricing_table: 2 to 4 plans")
        spn = 12 // n
        cta_h = 0.0
        if cta or aside:
            cta_h = TYPE["button"].size * TYPE["button"].lh / 72 + 2 * 15 / PX_PER_IN
        top = y + 0.28
        bottom = BOTTOM - (cta_h + 0.28 if cta_h else 0)
        padx, pady = 34 / PX_PER_IN, 26 / PX_PER_IN
        self._pricing_cards(s, plans, top, bottom, padx, pady, spn)
        if cta_h:
            bw = self.button_width(cta) if cta else 0
            aw = self.text_width(aside, "small") + 0.05 if aside else 0
            gap = 28 / PX_PER_IN if (cta and aside) else 0
            tw = bw + gap + aw
            if tw > CONTENT_W:
                raise TextFitError(f"slide {s.page}: CTA row is {tw:.2f} in wide")
            x0 = (SLIDE_W - tw) / 2
            yb = BOTTOM - cta_h
            if cta:
                self.button(s, x0, yb, cta, name="cta")
            if aside:
                ah = TYPE["small"].size * TYPE["small"].lh / 72
                self.text(s, x0 + bw + gap, yb + (cta_h - ah) / 2, aw, ah + 0.0005, aside, "small", name="cta aside")
        return s


    def _pricing_cards(self, s, plans, top, bottom, padx, pady, spn):
        """Cards first, then one feature-list scale shared by every card, then the lists."""
        pending = []
        for i, pl in enumerate(plans):
            x, w = col_x(1 + i * spn), span(spn)
            hl = bool(pl.get("highlight"))
            cardel = self.card(s, x, top, w, bottom - top, fill="white", radius=24 / PX_PER_IN,
                               line="red" if hl else "hairline", line_w=1.2 if hl else HAIR_PT,
                               name=f"plan card {pl['name']}")
            ix, iw = x + padx, w - 2 * padx
            yy = top + pady
            tag_w = 0.0
            if pl.get("tag"):
                tag_w = self.text_width(pl["tag"], "tag") + 0.05
                self.text(s, ix + iw - tag_w, yy + 0.07, tag_w, None, pl["tag"], "tag", align="r",
                          parent=cardel, name="plan tag")
            nm = self.text(s, ix, yy, iw - tag_w - (0.1 if tag_w else 0), None, pl["name"], "plan_name",
                           max_h=0.3, parent=cardel, name="plan name")
            yy = nm.b + 14 / PX_PER_IN
            pp = self.P(pl["price"], "price")
            if pl.get("unit"):
                pp.runs.append(Run("  " + pl["unit"], "regular", TYPE["unit"].size, col("ink2"), False, 0))
            pr = self.text(s, ix, yy, iw, None, [pp], "price", max_h=TYPE["price"].size * 1.12 / 72 + 0.001,
                           parent=cardel, name="plan price")
            yy = pr.b + 10 / PX_PER_IN
            if pl.get("inr"):
                ip = self.P(pl["inr"], "inr")
                if pl.get("inr_aside"):
                    ip.runs.append(Run(" · " + pl["inr_aside"], "regular", TYPE["inr"].size, col("ink3"), False, 0))
                el = self.text(s, ix, yy, iw, None, [ip], "inr", max_h=2 * 12 * 1.3 / 72 + 0.001,
                               parent=cardel, name="plan inr")
                yy = el.b + 6 / PX_PER_IN
            if pl.get("note"):
                el = self.text(s, ix, yy, iw, None, pl["note"], "note", max_h=2 * 9 * 1.4 / 72 + 0.001,
                               parent=cardel, name="plan note")
                yy = el.b
            yy += 16 / PX_PER_IN
            self.line(s, ix, yy, ix + iw, yy, parent=cardel, name="plan rule")
            yy += 16 / PX_PER_IN
            pending.append((pl, cardel, ix, iw, yy, bottom - pady - yy))

        def lay(pl, iw, k):
            fw = iw - 22 / PX_PER_IN
            ps = [self.P(f, "plan_list").scaled(k) for f in pl.get("features") or []]
            gap = 7 / PX_PER_IN * k
            hs = [self.measure([p], fw)[0] for p in ps]
            return ps, hs, gap, sum(hs) + gap * max(0, len(hs) - 1)

        k = 1.0
        while any(lay(pl, iw, k)[3] > avail + 1e-6 for pl, _, _, iw, _, avail in pending):
            k = round(k - 0.025, 3)
            if k < 0.9 - 1e-9:
                raise TextFitError(f"slide {s.page}: plan feature lists do not fit their cards")
        for pl, cardel, ix, iw, yy, avail in pending:
            ps, hs, gap, _ = lay(pl, iw, k)
            for p, hh in zip(ps, hs):
                self.mark(s, "check", ix, yy + 0.035 * k, 0.1, color="ink", width=1.1, parent=cardel,
                          name="feature check")
                self.text(s, ix + 22 / PX_PER_IN, yy, iw - 22 / PX_PER_IN, hh + 0.0005, [p], "plan_list",
                          parent=cardel, name="plan feature")
                yy += hh + gap

    # ------------------------------------------------------------------ timeline
    def timeline(self, title: str, nodes: Sequence[dict], *, eyebrow: Optional[str] = None,
                 lead: Optional[str] = None, notes: str = "", footer=None, page: Optional[str] = None,
                 hidden: bool = False, thumb_h: float = 1.15) -> SlideRec:
        """3-6 nodes on a hairline rail. node: {time, text, image (optional), active: bool}.
        Thumbnails sit in dark rounded cards above the rail; the active node gets a red dot."""
        s = self._new("timeline", "white", notes, footer, page, hidden)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        n = len(nodes)
        if not 3 <= n <= 6:
            raise LayoutError("timeline: 3 to 6 nodes")
        cw = (CONTENT_W - (n - 1) * GUTTER) / n
        xs = [MARGIN + i * (cw + GUTTER) for i in range(n)]
        has_img = any(nd.get("image") for nd in nodes)
        th = TYPE["h3"].size * TYPE["h3"].lh / 72 + 0.0005
        text_h = max(self.measure([self.P(nd["text"], "body")], cw)[0] for nd in nodes) + 0.0005
        block = (thumb_h + 0.35 if has_img else 0) + 0.25 + th + 0.1 + text_h
        top0 = y + 0.4
        if top0 + block > BOTTOM + 1e-6:
            raise LayoutError(f"timeline {s.page}: needs {block:.2f} in below the header")
        yy = top0 + (BOTTOM - top0 - block) / 2
        if has_img:
            for xi, nd in zip(xs, nodes):
                if not nd.get("image"):
                    continue
                cardel = self.rect(s, xi, yy, cw, thumb_h, grad=[(0, "scr_top"), (1, "scr_mid")],
                                   radius=12 / PX_PER_IN, name="timeline thumb card")
                path = self.image_path(nd["image"])
                f = self.fit_image(path, xi + 0.06, yy + 0.06, cw - 0.12, thumb_h - 0.12, mode="contain")
                self.picture(s, path, f["x"], f["y"], f["w"], f["h"], crop=f["crop"], radius=0.04,
                             parent=cardel, name=f"timeline thumb {nd['image']}")
            yy += thumb_h + 0.35
        rail = self.line(s, MARGIN, yy, MARGIN + CONTENT_W, yy, name="timeline rail")
        for xi, nd in zip(xs, nodes):
            d = 0.16 if nd.get("active") else 0.12
            self.rect(s, xi, yy - d / 2, d, d, oval=True, fill="red" if nd.get("active") else "ink",
                      hosts=(rail,), name="timeline dot")
            self.text(s, xi, yy + 0.25, cw, th, nd["time"], "h3", name="timeline time")
            self.text(s, xi, yy + 0.25 + th + 0.1, cw, text_h, nd["text"], "body", name="timeline text")
        return s

    # ------------------------------------------------------------------ CTA / quote
    def quote_or_cta(self, title: Optional[str] = None, *, eyebrow: Optional[str] = None,
                     lead: Optional[str] = None, button: Optional[str] = None, caption: Optional[str] = None,
                     extra: Sequence[str] = (), quote: Optional[str] = None, attribution: Optional[str] = None,
                     image: Optional[str] = None, notes: str = "", footer=None, page: Optional[str] = None,
                     hidden: bool = False, lid_w: float = 8.0) -> SlideRec:
        """Closing slide. CTA mode: centred headline, lead, red pill with a caption beside it,
        extra lines (e.g. [TBD] links) and an optional MacBook bleeding off the bottom.
        Quote mode (quote=...): large light quote and attribution, centred."""
        s = self._new("quote_or_cta", "mist", notes, footer if not image else False, page, hidden)
        x, w = MARGIN, CONTENT_W
        if quote:
            ts = TYPE["quote"]
            qh = self.measure([self.P(quote, "quote", align="c")], span(10))[0]
            ah = TYPE["body"].size * TYPE["body"].lh / 72 if attribution else 0
            blk = qh + (0.3 + ah if attribution else 0)
            yy = TOP + (BOTTOM - TOP - blk) / 2
            q = self.text(s, col_x(2), yy, span(10), None, quote, "quote", align="c",
                          max_h=BOTTOM - TOP - 0.6, name="quote")
            if attribution:
                self.text(s, col_x(2), q.b + 0.3, span(10), None, attribution, "body", align="c", name="attribution")
            return s
        y = self._header(s, eyebrow, title, lead, align="c", h1="h1_large", title_lines=2)
        if button or caption:
            bh = TYPE["button"].size * TYPE["button"].lh / 72 + 2 * 15 / PX_PER_IN
            bw = self.button_width(button) if button else 0
            cw_ = min(self.text_width(caption, "small") + 0.05, CONTENT_W - bw - 0.3) if caption else 0
            gap = 28 / PX_PER_IN if (button and caption) else 0
            x0 = (SLIDE_W - (bw + gap + cw_)) / 2
            yb = y + 0.3
            if button:
                self.button(s, x0, yb, button, name="cta")
            if caption:
                ch = self.measure([self.P(caption, "small")], cw_)[0]
                self.text(s, x0 + bw + gap, yb + max(0, (bh - ch) / 2), cw_, ch + 0.0005, caption, "small",
                          name="cta caption")
                bh = max(bh, ch)
            y = yb + bh
        for line_ in extra:
            el = self.text(s, x, y + 0.15, w, None, line_, "small", align="c", name="cta extra")
            y = el.b
        if image:
            dy = max(y + 0.4, 3.4)
            if dy + 1.2 > SLIDE_H:
                raise LayoutError(f"quote_or_cta {s.page}: no room for the device")
            self.macbook(s, (SLIDE_W - lid_w) / 2, dy, lid_w, image, pad_px=14, shot_frac=1.0,
                         clip_bottom=SLIDE_H, name="device")
        return s

    # ================================================================== output
    def save(self, path: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.prs.save(path)
        return path

    # ------------------------------------------------------------------ checks
    def check_layout(self) -> list:
        """Return a list of problems: dict(slide, severity, check, element, message)."""
        probs = []

        def add(s, sev, chk, el, msg):
            probs.append({"slide": s.page, "severity": sev, "check": chk,
                          "element": el.name if el else "", "message": msg})

        for s in self.slides:
            by_id = {e.eid: e for e in s.els}
            red_area = 0.0
            for e in s.els:
                # 1. text fit (independent re-measure)
                if e.kind == "text":
                    il, it, ir, ib = e.props["inset"]
                    need, ov, counts = self.measure(e.props["paras"], e.w - il - ir)
                    if need > e.h - it - ib + 0.5 / 72:
                        add(s, "error", "text-fit", e, f"needs {need:.3f} in, box {e.h - it - ib:.3f} in "
                                                       f"({sum(counts)} lines)")
                    if ov:
                        add(s, "error", "text-fit", e, f"word {ov!r} wider than the box")
                    for p in e.props["paras"]:
                        for r in p.runs:
                            if r.text.strip() and r.size < MIN_FONT_PT - 1e-6:
                                add(s, "error", "min-size", e, f"{r.size:.1f} pt < {MIN_FONT_PT} pt")
                            if r.color == col("red") and r.text.strip():
                                red_area += min(e.w, self.m.width(r.text, r.weight, r.italic, r.size) / 72) \
                                    * p.pitch / 72
                    if e.props.get("fill") == col("red"):
                        red_area += e.w * e.h
                elif e.kind == "rect":
                    if e.props.get("fill") == col("red"):
                        red_area += e.w * e.h
                    if e.props.get("line") == col("red"):
                        red_area += 2 * (e.w + e.h) * e.props["line_w"] / 72
                elif e.kind == "image":
                    iw, ih = e.props["px"]
                    l, t, r, b = e.props["crop"]
                    vis_w, vis_h = iw * (1 - l - r), ih * (1 - t - b)
                    up = e.w * self.ref_ppi / vis_w
                    if up > self.max_upscale + 1e-3:
                        add(s, "error", "image-upscale", e, f"shown at {up:.2f}x (max {self.max_upscale}x)")
                    if abs((e.w / e.h) / (vis_w / vis_h) - 1) > 0.01:
                        add(s, "error", "image-aspect", e, "aspect ratio distorted")
                elif e.kind == "chart":
                    sp = e.props["spec"]
                    if sp["type"] in ("column", "column_stacked", "line"):
                        slot = e.w / max(1, len(sp["categories"])) * 72 * 0.92
                        for c in sp["categories"]:
                            n, ov, _ = self.wrap(self.P(str(c), "small", size=9), slot)
                            if n > 2 or ov:
                                add(s, "warn", "chart-label", e, f"category label {c!r} needs {n} lines")
                # 2. containment / safe area
                if e.parent is not None:
                    p = by_id.get(e.parent)
                    if p and (e.x < p.x - EPS or e.y < p.y - EPS or e.r > p.r + EPS or e.b > p.b + EPS):
                        add(s, "error", "containment", e, f"outside its container {p.name} "
                                                          f"{e.box()} vs {p.box()}")
                else:
                    if e.bleed:
                        if e.x < -EPS or e.y < -EPS or e.r > SLIDE_W + EPS or e.b > SLIDE_H + EPS:
                            add(s, "error", "slide-bounds", e, f"off the slide {e.box()}")
                    elif e.x < SAFE[0] - EPS or e.y < SAFE[1] - EPS or e.r > SAFE[2] + EPS or e.b > SAFE[3] + EPS:
                        add(s, "error", "safe-area", e, f"outside the safe area {e.box()}")
            # 3. overlaps between siblings that are not intentional layers
            groups: dict = {}
            for e in s.els:
                groups.setdefault(e.parent, []).append(e)
            for sibs in groups.values():
                for i in range(len(sibs)):
                    a = sibs[i]
                    for b in sibs[i + 1:]:
                        if a.eid in b.hosts or b.eid in a.hosts:
                            continue
                        if _overlap(a, b):
                            add(s, "error", "overlap", a, f"overlaps {b.name}: {a.box()} x {b.box()}")
            share = red_area / (SLIDE_W * SLIDE_H)
            if share > 0.10:
                add(s, "warn", "red-share", None, f"red covers about {share:.1%} of the slide (keep under 10%)")
        return probs

    # ------------------------------------------------------------------ HTML preview
    def html_preview(self, path: str, *, problems: Optional[list] = None) -> str:
        """Write an HTML preview: one 1600 x 900 frame per slide, built from the registry.
        index.html?slide=N shows slide N alone at 1:1 (for headless screenshots)."""
        out_dir = os.path.dirname(os.path.abspath(path))
        img_dir = os.path.join(out_dir, "img")
        os.makedirs(img_dir, exist_ok=True)
        problems = self.check_layout() if problems is None else problems
        copied = {}

        def img_url(src):
            if src not in copied:
                base = os.path.basename(src)
                dst = os.path.join(img_dir, base)
                if os.path.abspath(src) != os.path.abspath(dst):
                    shutil.copyfile(src, dst)
                copied[src] = "img/" + base
            return copied[src]

        weights = {k: self.fp[k][3] for k in ("light", "regular", "medium")}
        frames = []
        for s in self.slides:
            parts = []
            for e in s.els:
                parts.append(self._html_el(e, weights, img_url))
            sp = [p for p in problems if p["slide"] == s.page]
            prob_html = ""
            if sp:
                prob_html = "<ul class='probs'>" + "".join(
                    f"<li class='{p['severity']}'>{_html.escape(p['check'])}: {_html.escape(p['element'])} "
                    f"{_html.escape(p['message'])}</li>" for p in sp) + "</ul>"
            frames.append(
                f"<section class='sl' id='s{s.index}' data-page='{_html.escape(s.page)}'>"
                f"<div class='lbl'>{_html.escape(s.page)} · {s.layout}{' · hidden' if s.hidden else ''}</div>"
                f"<div class='slot'><div class='frame' style='background:{s.bg}'>{''.join(parts)}</div></div>"
                f"{prob_html}"
                + (f"<details class='notes'><summary>Speaker notes</summary><p>{_html.escape(s.notes)}</p></details>"
                   if s.notes else "")
                + "</section>")
        n_err = sum(1 for p in problems if p["severity"] == "error")
        n_warn = sum(1 for p in problems if p["severity"] == "warn")
        doc = _HTML_TEMPLATE.format(
            title=_html.escape(self.title), font=self.fp["css"], link=self.fp["css_link"],
            frames="\n".join(frames), n=len(self.slides), errs=n_err, warns=n_warn,
            profile=self.profile_name)
        with open(path, "w", encoding="utf-8") as f:
            f.write(doc)
        return path

    def _html_el(self, e: El, weights: dict, img_url) -> str:
        box = f"left:{_px(e.x)}px;top:{_px(e.y)}px;width:{_px(e.w)}px;height:{_px(e.h)}px"
        if e.kind == "rect":
            p = e.props
            st = [box]
            if p["grad"]:
                st.append("background:linear-gradient(to bottom," +
                          ",".join(f"{c} {pos * 100:.2f}%" for pos, c in p["grad"]) + ")")
            elif p["fill"]:
                st.append(f"background:{p['fill']}")
            if p["line"]:
                lw = p["line_w"] / PT_PER_PX
                st.append(f"box-shadow:inset 0 0 0 {lw:.2f}px {p['line']}")
            if p["oval"]:
                st.append("border-radius:50%")
            else:
                tl, tr, br, bl = (_px(r) for r in p["radii"])
                if any((tl, tr, br, bl)):
                    st.append(f"border-radius:{tl}px {tr}px {br}px {bl}px")
            return f"<div class='r' data-name='{_html.escape(e.name)}' style='{';'.join(st)}'></div>"
        if e.kind == "line":
            x1, y1, x2, y2 = e.props["pts"]
            lw = e.props["width"] / PT_PER_PX
            if abs(y1 - y2) < 1e-9:
                return (f"<div class='ln' style='left:{_px(min(x1, x2))}px;top:{_px(y1) - lw / 2:.2f}px;"
                        f"width:{_px(abs(x2 - x1))}px;height:{lw:.2f}px;background:{e.props['color']}'></div>")
            if abs(x1 - x2) < 1e-9:
                return (f"<div class='ln' style='left:{_px(x1) - lw / 2:.2f}px;top:{_px(min(y1, y2))}px;"
                        f"width:{lw:.2f}px;height:{_px(abs(y2 - y1))}px;background:{e.props['color']}'></div>")
            return (f"<svg class='ln' style='left:0;top:0' width='1600' height='900'><line x1='{_px(x1)}' "
                    f"y1='{_px(y1)}' x2='{_px(x2)}' y2='{_px(y2)}' stroke='{e.props['color']}' "
                    f"stroke-width='{lw:.2f}'/></svg>")
        if e.kind == "mark":
            sz = _px(e.w)
            lw = e.props["width"] / PT_PER_PX
            segs = "".join(f"<line x1='{a[0] * sz:.2f}' y1='{a[1] * sz:.2f}' x2='{b[0] * sz:.2f}' y2='{b[1] * sz:.2f}'/>"
                           for a, b in e.props["segs"])
            return (f"<svg class='mk' style='{box}' viewBox='0 0 {sz} {sz}' stroke='{e.props['color']}' "
                    f"stroke-width='{lw:.2f}' stroke-linecap='round'>{segs}</svg>")
        if e.kind == "image":
            l, t, r, b = e.props["crop"]
            fw = e.w / max(1e-6, 1 - l - r)
            fh = e.h / max(1e-6, 1 - t - b)
            rad = f"border-radius:{_px(e.props['radius'])}px;" if e.props["radius"] else ""
            return (f"<div class='im' style='{box};{rad}'><img alt='' src='{img_url(e.props['src'])}' "
                    f"style='left:{-_px(fw * l)}px;top:{-_px(fh * t)}px;width:{_px(fw)}px;height:{_px(fh)}px'></div>")
        if e.kind == "text":
            p = e.props
            il, it, ir, ib = p["inset"]
            jc = {"t": "flex-start", "m": "center", "b": "flex-end"}[p["anchor"]]
            st = [box, f"justify-content:{jc}",
                  f"padding:{_px(it)}px {_px(ir)}px {_px(ib)}px {_px(il)}px"]
            if p.get("fill"):
                st.append(f"background:{p['fill']}")
                if p.get("radius"):
                    st.append(f"border-radius:{_px(p['radius'])}px")
            ps = []
            for i, para in enumerate(p["paras"]):
                runs = "".join(
                    f"<span style='font-weight:{weights[r.weight]};font-size:{r.size / PT_PER_PX:.2f}px;color:{r.color};"
                    f"{'font-style:italic;' if r.italic else ''}"
                    f"{f'letter-spacing:{r.tracking}em;' if r.tracking else ''}'>"
                    f"{_html.escape(r.text).replace(chr(10), '<br>')}</span>" for r in para.runs)
                ps.append(f"<p style='text-align:{ {'l': 'left', 'c': 'center', 'r': 'right'}[para.align]};"
                          f"line-height:{para.pitch / PT_PER_PX:.2f}px;font-size:{para.size / PT_PER_PX:.2f}px;"
                          f"margin-top:{(para.space_before / PT_PER_PX) if i else 0:.2f}px'>{runs}</p>")
            return f"<div class='t' data-name='{_html.escape(e.name)}' style='{';'.join(st)}'>{''.join(ps)}</div>"
        if e.kind == "chart":
            return f"<div class='ch' style='{box}'>{self._chart_svg(e.props['spec'], _px(e.w), _px(e.h))}</div>"
        return ""

    # ------------------------------------------------------------------ chart preview (SVG)
    def _chart_svg(self, sp: dict, W: float, H: float) -> str:
        fs = 9 / PT_PER_PX             # 15 px
        ink, ink2, hair = col("ink"), col("ink2"), col("hairline")
        wmed = self.fp["medium"][3]
        out = [f"<svg width='{W}' height='{H}' viewBox='0 0 {W} {H}' font-size='{fs:.1f}'>"]
        t = sp["type"]
        cats, series = sp["categories"], sp["series"]
        fmt = sp["number_format"]

        def txt(x, y, s_, anchor="middle", color=ink2, weight=400):
            return (f"<text x='{x:.1f}' y='{y:.1f}' text-anchor='{anchor}' fill='{color}' "
                    f"font-weight='{weight}'>{_html.escape(str(s_))}</text>")

        if t == "doughnut":
            ser = series[0]
            vals = [v or 0 for v in ser["values"]]
            tot = sum(vals) or 1
            cx, cy, R = W / 2, H / 2, min(W, H) / 2 - 4
            r0 = R * 0.62
            a = -math.pi / 2
            for i, v in enumerate(vals):
                if v <= 0:
                    continue
                b = a + 2 * math.pi * v / tot
                large = 1 if b - a > math.pi else 0
                if v >= tot:
                    b = a + 2 * math.pi * 0.9999
                p = (f"M{cx + R * math.cos(a):.2f},{cy + R * math.sin(a):.2f} "
                     f"A{R},{R} 0 {large} 1 {cx + R * math.cos(b):.2f},{cy + R * math.sin(b):.2f} "
                     f"L{cx + r0 * math.cos(b):.2f},{cy + r0 * math.sin(b):.2f} "
                     f"A{r0},{r0} 0 {large} 0 {cx + r0 * math.cos(a):.2f},{cy + r0 * math.sin(a):.2f}Z")
                out.append(f"<path d='{p}' fill='{ser['point_colors'][i]}' stroke='#fff' stroke-width='2.5'/>")
                a = b
            out.append("</svg>")
            return "".join(out)

        legend_h = 30 if sp["legend"] else 0
        lo, hi = sp["axis_min"], sp["axis_max"]
        horiz = t.startswith("bar")
        stacked = t.endswith("stacked")
        ns = len(series)
        if horiz:
            lab_w = min(W * 0.32, max(self.m.width(str(c), "regular", False, 9) / PT_PER_PX for c in cats) + 12)
            px0, px1 = lab_w + 6, W - 70
            py0, py1 = 4, H - legend_h - 4
            slot = (py1 - py0) / len(cats)
            scale = lambda v: px0 + (v - lo) / (hi - lo) * (px1 - px0)
            out.append(f"<line x1='{scale(0):.1f}' y1='{py0}' x2='{scale(0):.1f}' y2='{py1}' stroke='{hair}' stroke-width='1.25'/>")
            for ci, c in enumerate(cats):
                y0 = py0 + ci * slot
                out.append(txt(lab_w, y0 + slot / 2 + fs * 0.35, c, "end"))
                nb = 1 if stacked else ns
                bh = slot / (nb + 0.7)
                acc = 0.0
                for si, ser in enumerate(series):
                    v = ser["values"][ci]
                    if v is None:
                        continue
                    clr = ser["point_colors"].get(ci, ser["color"])
                    if stacked:
                        xa, xb = scale(acc), scale(acc + v)
                        acc += v
                        by = y0 + 0.35 * bh
                    else:
                        xa, xb = scale(0), scale(v)
                        by = y0 + 0.35 * bh + si * bh
                    out.append(f"<rect x='{min(xa, xb):.1f}' y='{by:.1f}' width='{abs(xb - xa):.1f}' height='{bh:.1f}' fill='{clr}'/>")
                    if stacked:
                        lx, anc = (xa + xb) / 2, "middle"
                    elif v >= 0:
                        lx, anc = max(xa, xb) + 6, "start"
                    else:
                        lx, anc = min(xa, xb) - 6, "end"
                    out.append(txt(lx, by + bh / 2 + fs * 0.35, fmt_num(v, fmt), anc, ink, wmed))
        else:
            cat_lines = []
            slot_w = W / len(cats)
            for c in cats:
                _, _, lines = self.wrap(self.P(str(c), "small", size=9), slot_w * 0.92 * PT_PER_PX * 1.0)
                cat_lines.append(lines[:2])
            lab_h = max(len(l_) for l_ in cat_lines) * fs * 1.25 + 8
            px0, px1 = 4, W - 4
            py0, py1 = fs * 1.8, H - legend_h - lab_h
            slot = (px1 - px0) / len(cats)
            scale = lambda v: py1 - (v - lo) / (hi - lo) * (py1 - py0)
            out.append(f"<line x1='{px0}' y1='{scale(0):.1f}' x2='{px1}' y2='{scale(0):.1f}' stroke='{hair}' stroke-width='1.25'/>")
            for ci, lines in enumerate(cat_lines):
                for li, l_ in enumerate(lines):
                    out.append(txt(px0 + ci * slot + slot / 2, py1 + 8 + fs + li * fs * 1.25, l_))
            if t == "line":
                for si, ser in enumerate(series):
                    pts = [(px0 + ci * slot + slot / 2, scale(v) if v is not None else 0, v) for ci, v in enumerate(ser["values"])]
                    seg = []
                    for p_ in pts + [(0, 0, None)]:
                        if p_[2] is None:
                            if len(seg) > 1:
                                out.append("<polyline fill='none' stroke='{}' stroke-width='3.75' points='{}'/>".format(
                                    ser["color"], " ".join(f"{a:.1f},{b:.1f}" for a, b, _ in seg)))
                            seg = []
                        else:
                            seg.append(p_)
                    for ci, (a, b, v) in enumerate(pts):
                        if v is None:
                            continue
                        clr = ser["point_colors"].get(ci, ser["color"])
                        out.append(f"<circle cx='{a:.1f}' cy='{b:.1f}' r='5.8' fill='{clr}'/>")
                        out.append(txt(a, b - 12, fmt_num(v, fmt), "middle", ink, wmed))
            else:
                nb = 1 if stacked else ns
                gapw = 0.7
                bw = slot / (nb + gapw)
                for ci in range(len(cats)):
                    x0 = px0 + ci * slot + gapw * bw / 2
                    acc = 0.0
                    for si, ser in enumerate(series):
                        v = ser["values"][ci]
                        if v is None:
                            continue
                        clr = ser["point_colors"].get(ci, ser["color"])
                        if stacked:
                            ya, yb = scale(acc), scale(acc + v)
                            acc += v
                            bx = x0
                        else:
                            ya, yb = scale(0), scale(v)
                            bx = x0 + si * bw
                        out.append(f"<rect x='{bx:.1f}' y='{min(ya, yb):.1f}' width='{bw * (0.94 if not stacked and ns > 1 else 1):.1f}' "
                                   f"height='{abs(yb - ya):.1f}' fill='{clr}'/>")
                        ly = (ya + yb) / 2 + fs * 0.35 if stacked else min(ya, yb) - 6
                        out.append(txt(bx + bw / 2, ly, fmt_num(v, fmt), "middle", ink, wmed))
        if sp["legend"]:
            items = [(ser["name"], ser["color"]) for ser in series]
            widths = [self.m.width(n_, "regular", False, 9) / PT_PER_PX + 30 for n_, _ in items]
            x = (W - sum(widths)) / 2
            yl = H - legend_h / 2
            for (n_, c_), w_ in zip(items, widths):
                out.append(f"<rect x='{x:.1f}' y='{yl - 6:.1f}' width='12' height='12' fill='{c_}'/>")
                out.append(txt(x + 18, yl + fs * 0.35, n_, "start"))
                x += w_
        out.append("</svg>")
        return "".join(out)


def _overlap(a: El, b: El) -> bool:
    """True if the boxes overlap by more than EPS. Zero-thickness rules count only when they
    run through the inside of the other box."""
    def axis(a0, a1, b0, b1):
        if a1 - a0 < EPS and b1 - b0 < EPS:
            return abs(a0 - b0) < EPS
        if a1 - a0 < EPS:
            return b0 + EPS < a0 < b1 - EPS
        if b1 - b0 < EPS:
            return a0 + EPS < b0 < a1 - EPS
        return min(a1, b1) - max(a0, b0) > EPS
    return axis(a.x, a.r, b.x, b.r) and axis(a.y, a.b, b.y, b.b)


def chrome_check(index_html: str, chrome: str = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
                 timeout: int = 60) -> list:
    """Render the preview in headless Chrome and return text boxes whose content overflows
    (browser-measured, a second opinion on the PIL estimate)."""
    if not os.path.exists(chrome):
        return [{"severity": "warn", "check": "chrome", "message": "Chrome not found; browser check skipped"}]
    url = "file://" + os.path.abspath(index_html) + "?qa=1"
    r = subprocess.run([chrome, "--headless=new", "--disable-gpu", "--virtual-time-budget=6000",
                        "--window-size=1600,900", "--dump-dom", url],
                       capture_output=True, text=True, timeout=timeout)
    m = re.search(r"<pre id=\"qa\">(.*?)</pre>", r.stdout, re.S)
    if not m:
        return [{"severity": "warn", "check": "chrome", "message": "no QA report in DOM"}]
    return json.loads(_html.unescape(m.group(1)))


def chrome_screenshot(index_html: str, slide: int, out_png: str,
                      chrome: str = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome") -> str:
    """Screenshot one preview frame at 1600 x 900 with headless Chrome."""
    url = "file://" + os.path.abspath(index_html) + f"?slide={slide}"
    subprocess.run([chrome, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--virtual-time-budget=4000",
                    "--window-size=1600,900", f"--screenshot={out_png}", url],
                   capture_output=True, text=True, timeout=60)
    return out_png


_HTML_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title} preview</title>
{link}
<style>
  :root {{ --page:#E8E8ED; --ink:#1D1D1F; --ink2:#6E6E73; }}
  * {{ box-sizing:border-box; margin:0; padding:0; }}
  body {{ background:var(--page); color:var(--ink); font-family:{font}; -webkit-font-smoothing:antialiased;
         padding:32px 16px 64px; }}
  header.top {{ max-width:1600px; margin:0 auto 24px; }}
  header.top h1 {{ font-size:24px; font-weight:500; }}
  header.top p {{ color:var(--ink2); font-size:14px; margin-top:6px; }}
  .sl {{ max-width:1600px; margin:0 auto 36px; }}
  .lbl {{ font-size:13px; color:var(--ink2); margin-bottom:8px; }}
  .slot {{ position:relative; width:100%; aspect-ratio:16/9; overflow:hidden; border-radius:10px;
          box-shadow:0 1px 2px rgba(0,0,0,.06),0 8px 28px rgba(0,0,0,.06); }}
  .frame {{ position:absolute; left:0; top:0; width:1600px; height:900px; overflow:hidden;
           transform-origin:0 0; transform:scale(var(--k,1)); font-family:{font}; }}
  .frame > * {{ position:absolute; }}
  .t {{ display:flex; flex-direction:column; overflow:visible; }}
  .t p {{ margin:0; white-space:normal; overflow-wrap:normal; }}
  .t p span {{ line-height:0; }}
  .im {{ overflow:hidden; }}
  .im img {{ position:absolute; display:block; max-width:none; }}
  .ln {{ pointer-events:none; }}
  .mk {{ overflow:visible; }}
  .ch svg {{ display:block; overflow:visible; font-family:{font}; }}
  .probs {{ margin:8px 0 0 18px; font-size:13px; }}
  .probs .error {{ color:#C40812; }} .probs .warn {{ color:#8A6D00; }}
  .notes {{ margin-top:8px; font-size:14px; color:var(--ink2); max-width:1100px; line-height:1.5; }}
  .notes summary {{ cursor:pointer; color:var(--ink); }}
  body.one {{ padding:0; background:#fff; }}
  body.one header.top, body.one .lbl, body.one .probs, body.one .notes {{ display:none; }}
  body.one .sl {{ display:none; margin:0; }} body.one .sl.on {{ display:block; }}
  body.one .slot {{ width:1600px; height:900px; border-radius:0; box-shadow:none; }}
  #qa {{ display:none; }}
</style>
</head>
<body>
<header class="top">
  <h1>{title}</h1>
  <p>{n} slides · font profile {profile} · layout check: {errs} errors, {warns} warnings. Frames are 1600 × 900 (120 px per inch), built from the same layout data as the .pptx.</p>
</header>
{frames}
<pre id="qa"></pre>
<script>
(function () {{
  var q = new URLSearchParams(location.search);
  var one = q.get('slide');
  if (one) {{
    document.body.classList.add('one');
    var el = document.getElementById('s' + one); if (el) el.classList.add('on');
  }}
  function fit() {{
    document.querySelectorAll('.slot').forEach(function (s) {{
      var k = one ? 1 : s.clientWidth / 1600;
      s.firstElementChild.style.setProperty('--k', k);
    }});
  }}
  fit();
  window.addEventListener('resize', fit);
  function qa() {{
    var out = [];
    document.querySelectorAll('.t').forEach(function (t) {{
      var sl = t.closest('.sl');
      var cs = getComputedStyle(t);
      var inner = t.clientHeight - parseFloat(cs.paddingTop) - parseFloat(cs.paddingBottom);
      var need = 0;
      Array.prototype.forEach.call(t.children, function (p) {{
        need += p.offsetHeight + parseFloat(getComputedStyle(p).marginTop);
        if (p.scrollWidth > p.clientWidth + 1) out.push({{slide: sl.dataset.page, severity: 'error', check: 'browser-width', element: t.dataset.name, message: 'line wider than box'}});
      }});
      if (need > inner + 1.5) out.push({{slide: sl.dataset.page, severity: 'error', check: 'browser-fit', element: t.dataset.name, message: 'content ' + need.toFixed(1) + 'px > box ' + inner.toFixed(1) + 'px'}});
    }});
    document.getElementById('qa').textContent = JSON.stringify(out);
  }}
  if (q.get('qa')) {{
    (document.fonts && document.fonts.ready ? document.fonts.ready : Promise.resolve()).then(qa);
  }}
}})();
</script>
</body>
</html>
"""
