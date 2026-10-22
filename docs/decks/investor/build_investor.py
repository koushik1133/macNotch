"""
build_investor.py - SparkV seed investor deck, Apple-minimal style.

Reads   docs/decks/investor/outline.json  (copy, notes, chart data, image names)
Uses    docs/decks/generator/sparkv_deck.py (layouts, checks, HTML preview; not modified)
Writes  docs/decks/investor/SparkV-Investor-Deck.pptx
        docs/decks/investor/preview/index.html   (1600 x 900 frames, same layout data)
        docs/decks/investor/layout-report.json

Run:    /opt/anaconda3/bin/python3 build_investor.py [--chrome] [--shots DIR]

All headlines, body lines, numbers, notes, chart data and [TBD] text come from the outline.
Dense outline slides are split into two slides (STYLE.md section 7: one idea per slide) instead
of shrinking type: 8, 9, 11, 12 and 14 become two slides each, and the sources appendix runs
over as many hidden slides as it needs. Layouts the library lacks (waterfall, positioning map,
lanes, chips, image grid, matrix with a highlighted column) are built here from the library's
primitives, so the layout check and the preview cover them too.
"""
import argparse
import json
import math
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
DECKS = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(DECKS, "generator"))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_LABEL_POSITION  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Pt  # noqa: E402

import sparkv_deck as sd  # noqa: E402
from sparkv_deck import (BOTTOM, CONTENT_W, GUTTER, MARGIN, PX_PER_IN, SLIDE_W, TOP, TYPE,  # noqa: E402
                         Deck, LayoutError, TextFitError, accent, chrome_check, chrome_screenshot,
                         col, col_x, fmt_num, load_outline, span)

OUT_PPTX = os.path.join(HERE, "SparkV-Investor-Deck.pptx")
OUT_PREVIEW = os.path.join(HERE, "preview", "index.html")
OUT_REPORT = os.path.join(HERE, "layout-report.json")
PX = 1 / PX_PER_IN
OK_GREEN = "#1D8E3E"          # apple_minimal.ok in outline.styles

# Semibold emphasis run per headline (Apple minimal: Light headline with one Semibold phrase).
# Every phrase is a substring of the outline header; no copy changes.
EMPH = {
    "title": "AI command center",
    "problem": "rent five apps",
    "solution": "Every tool.",
    "live_demo": "all on the Mac",
    "meetings": "without a bot",
    "proof_ui_audit": "one real UI Audit run",
    "why_now": "already ships inside the Mac",
    "market": "paying, AI-curious",
    "competition": "one price",
    "business_model": "$39 once",
    "unit_economics": "about 94%",
    "go_to_market": "Product Hunt in 30 days",
    "traction": "Launch numbers come next",
    "roadmap": "all scheduled",
}


# Sample-data hygiene (BLUEPRINT rule 8): these renders show a sample note "Free, Pro $9, Team $19
# per seat" that contradicts the real pricing on the business-model slide. Until re-captures exist,
# use a clean render of the same feature from the outline's own image list.
SAFE_IMAGE = {
    "16-universal-search-panel.png": "15-universal-search.png",
    "36-meeting-notes.png": "37-meeting-live-answers.png",
}


def H(s):
    return accent(s["header"], EMPH.get(s["id"]))


def split_notes(text, k):
    """Split speaker notes after sentence k (1-based): (first k sentences, rest)."""
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z\[])", text.strip())
    return " ".join(parts[:k]), " ".join(parts[k:])


# =========================================================================== deck
class InvestorDeck(Deck):
    """Deck plus the investor-only layouts. Nothing in the library is changed."""

    # ------------------------------------------------------------------ primitives
    def dashed(self, s, el, color="ink3", width=0.75):
        """Give the shape behind registry element `el` (the last shape added) a dashed outline."""
        shp = s.slide.shapes[-1]
        shp.line.color.rgb = RGBColor.from_string(col(color)[1:])
        shp.line.width = Pt(width)
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        el.props["line"] = col(color)
        el.props["line_w"] = width
        el.props["dash"] = True
        return el

    def chip(self, s, x, y, text, *, style="small_ink", line="hairline", dashed=False, fill=None,
             parent=None, h=None, name="chip"):
        """Pill with an outline (or dashed outline) and centred one-line text. Returns the pill El."""
        padx = 16 * PX
        ts = TYPE[style]
        lh = ts.size * ts.lh / 72
        h = h or lh + 2 * 9 * PX
        w = self.text_width(text, style) + 2 * padx + 0.04
        pill = self.rect(s, x, y, w, h, fill=fill, line=None if dashed else line, radius=h / 2,
                         parent=parent, name=name)
        if dashed:
            self.dashed(s, pill)
        self.text(s, x + padx, y + (h - lh) / 2, w - 2 * padx, lh + 0.0005, text, style, align="c",
                  parent=pill, name=f"{name} text")
        return pill

    def chip_width(self, text, style="small_ink"):
        return self.text_width(text, style) + 2 * 16 * PX + 0.04

    def rule_rows(self, s, x, y, w, items, style="body", *, gap=14 * PX, max_bottom=BOTTOM, name="row"):
        """Hairline, text, hairline, text ... Returns the bottom y."""
        yy, last = y, y
        for it in items:
            self.line(s, x, yy, x + w, yy, name=f"{name} rule")
            el = self.text(s, x, yy + gap, w, None, it, style, max_h=max_bottom - yy - gap, name=name)
            last = el.b
            yy = el.b + gap
        if last > max_bottom + 1e-6:
            raise LayoutError(f"rows overflow: {last:.3f} > {max_bottom:.3f}")
        return last

    def para_h(self, text, style, w):
        return self.measure([self.P(text, style)], w)[0]

    # ------------------------------------------------------------------ chart card with notes
    def card_chart(self, s, x, y, w, h, spec, *, notes=(), title_color=None, dashed=False,
                   waterfall=None, legend_rows=None, fill="mist", name="chart card"):
        """Mist card: title, native chart (or waterfall / doughnut with legend rows),
        annotation lines (small, ink2), [TBD] lines and the source line under the chart."""
        sp = self._norm_chart(spec)
        cardel = self.card(s, x, y, w, h, fill=fill if not dashed else "white", name=name)
        if dashed:
            self.dashed(s, cardel)
        padx, pady = 32 * PX, 28 * PX
        ix, iw = x + padx, w - 2 * padx
        yy = y + pady
        if sp.get("title"):
            t = self.text(s, ix, yy, iw, None, sp["title"], "small_ink", color=title_color,
                          max_h=2 * 10.2 * 1.45 / 72 + 0.001, parent=cardel, name="chart title")
            yy = t.b + 0.12
        bottom = y + h - pady
        foot = [self.P(n_, "sub") for n_ in (sp.get("tbd") or [])]
        if sp.get("source"):
            src = sp["source"]
            foot.append(self.P(src if src.lower().startswith("source") else "Source: " + src, "sub"))
        if foot:
            fh = self.measure(foot, iw)[0]
            self.text(s, ix, bottom - fh - 0.0005, iw, fh + 0.0005, foot, "sub", parent=cardel, name="chart source")
            bottom -= fh + 0.1
        if notes:
            nps = [self.P(n_, "small") for n_ in notes]
            nh = self.measure(nps, iw)[0]
            self.text(s, ix, bottom - nh - 0.0005, iw, nh + 0.0005, nps, "small", parent=cardel,
                      name="chart annotation")
            bottom -= nh + 0.12
        if legend_rows:
            rh = 0.3
            lh = rh * len(legend_rows)
            ly = bottom - lh
            for i, (label, value, clr) in enumerate(legend_rows):
                ry = ly + i * rh
                self.rect(s, ix, ry + 0.07, 0.12, 0.12, fill=clr, radius=0.02, parent=cardel, name="legend swatch")
                self.text(s, ix + 0.22, ry, iw - 0.22 - 1.05, rh - 0.02, label, "small", parent=cardel,
                          name="legend label")
                self.text(s, ix + iw - 1.0, ry, 1.0, rh - 0.02, value, "small_ink", align="r", parent=cardel,
                          name="legend value")
            bottom = ly - 0.12
        ch_h = bottom - yy
        if ch_h < 1.2:
            raise LayoutError(f"slide {s.page}: chart area only {ch_h:.2f} in tall")
        if waterfall:
            self._waterfall(s, ix, yy, iw, ch_h, sp, waterfall, cardel)
        elif sp["type"] == "doughnut":
            d = min(ch_h, iw)
            self._native_chart(s, ix + (iw - d) / 2, yy, d, d, sp, self.fp["regular"][0], cardel)
        else:
            self.chart(s, ix, yy, iw, ch_h, sp, parent=cardel)
        return cardel

    def _waterfall(self, s, x, y, w, h, sp, wf, parent):
        """Waterfall as a stacked column chart: invisible base, visible step, invisible cap
        series whose centred custom label sits just above each step (python-pptx has no native
        waterfall chart)."""
        el = self._native_chart(s, x, y, w, h, sp, self.fp["regular"][0], parent)
        el.props["spec"]["waterfall"] = wf
        gf = s.slide.shapes[-1]
        plot = gf.chart.plots[0]
        base, step, cap = plot.series
        for ser in (base, cap):
            ser.format.fill.background()
            ser.format.line.fill.background()
        for ser in (base, step):
            dls = ser._element.get_or_add_dLbls()
            for ch in list(dls):
                dls.remove(ch)
            etree.SubElement(dls, qn("c:delete")).set("val", "1")
        for i, lab in enumerate(wf["labels"]):
            dl = cap.points[i].data_label
            dl.position = XL_LABEL_POSITION.CENTER
            tf = dl.text_frame
            tf.text = lab
            f = tf.paragraphs[0].runs[0].font
            f.size = Pt(9)
            f.name = self.fp["medium"][0]
            f.color.rgb = RGBColor.from_string(col(wf["label_colors"][i])[1:])
        return el

    # ------------------------------------------------------------------ HTML preview extras
    def _html_el(self, e, weights, img_url):
        out = super()._html_el(e, weights, img_url)
        if e.kind == "rect" and e.props.get("dash"):
            lw = e.props["line_w"] / sd.PT_PER_PX
            out = re.sub(r"box-shadow:inset 0 0 0 [0-9.]+px #[0-9A-F]{6}",
                         f"border:{lw:.2f}px dashed {e.props['line']}", out)
        return out

    def _chart_svg(self, sp, W, H):
        wf = sp.get("waterfall")
        if not wf:
            return super()._chart_svg(sp, W, H)
        fs = 9 / sd.PT_PER_PX
        ink2, hair = col("ink2"), col("hairline")
        wmed = self.fp["medium"][3]
        cats = sp["categories"]
        lo, hi = sp["axis_min"], sp["axis_max"]
        out = [f"<svg width='{W}' height='{H}' viewBox='0 0 {W} {H}' font-size='{fs:.1f}'>"]
        lab_h = fs * 1.25 + 8
        py0, py1 = 4, H - lab_h
        slot = W / len(cats)
        scale = lambda v: py1 - (v - lo) / (hi - lo) * (py1 - py0)  # noqa: E731
        out.append(f"<line x1='0' y1='{py1:.1f}' x2='{W}' y2='{py1:.1f}' stroke='{hair}' stroke-width='1.25'/>")
        bw = slot / 1.7
        for i, c in enumerate(cats):
            b, v, cp = wf["base"][i], wf["step"][i], wf["cap"]
            x0 = i * slot + (slot - bw) / 2
            if v > 0:
                out.append(f"<rect x='{x0:.1f}' y='{scale(b + v):.1f}' width='{bw:.1f}' "
                           f"height='{scale(b) - scale(b + v):.1f}' fill='{col(wf['colors'][i])}'/>")
            ly = (scale(b + v) + scale(b + v + cp)) / 2 + fs * 0.35
            out.append(f"<text x='{x0 + bw / 2:.1f}' y='{ly:.1f}' text-anchor='middle' "
                       f"fill='{col(wf['label_colors'][i])}' font-weight='{wmed}'>{wf['labels'][i]}</text>")
            out.append(f"<text x='{i * slot + slot / 2:.1f}' y='{py1 + 8 + fs:.1f}' text-anchor='middle' "
                       f"fill='{ink2}'>{c}</text>")
        out.append("</svg>")
        return "".join(out)

    # ------------------------------------------------------------------ title
    def investor_title(self, title, *, eyebrow, lead, line, image, notes, lid_w=9.167):
        s = self._new("title_slide", "mist", notes, False, None, False)
        e = self.text(s, MARGIN, TOP, CONTENT_W, None, eyebrow, "eyebrow", align="c", color="red", name="eyebrow")
        ts = TYPE["display"]
        t = self.text(s, MARGIN, e.b + 18 * PX, CONTENT_W, None, title, "display", align="c",
                      max_h=2 * ts.size * ts.lh / 72 + 0.001, name="headline")
        l_ = self.text(s, MARGIN, t.b + 26 * PX, CONTENT_W, None, lead, "lead", align="c",
                       max_h=2 * TYPE["lead"].size * TYPE["lead"].lh / 72 + 0.001, name="lead")
        c = self.text(s, MARGIN, l_.b + 14 * PX, CONTENT_W, None, line, "small", align="c", name="round line")
        dy = max(3.767, c.b + 0.25)
        self.macbook(s, (SLIDE_W - lid_w) / 2, dy, lid_w, image, pad_px=16, clip_bottom=sd.SLIDE_H,
                     name="hero device")
        return s

    # ------------------------------------------------------------------ problem (chart + tiles)
    def chart_and_tiles(self, title, *, eyebrow, lead, chart, metrics, after=None, notes="", footer=None,
                        chart_notes=(), tiles_side="right"):
        s = self._new("metrics_dashboard", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1_dash", title_lines=2, lead_lines=2)
        top = y + 0.3
        if tiles_side == "right":
            tx, tw, cx, cw = col_x(9), span(4), col_x(1), span(8)
        else:
            tx, tw, cx, cw = col_x(1), span(4), col_x(5), span(8)
        tb = BOTTOM
        if after:
            ah = self.para_h(after, "body_ink", tw)
            tb = BOTTOM - ah - 0.15 - 0.25
            self.line(s, tx, tb + 0.25, tx + tw, tb + 0.25, name="after rule")
            self.text(s, tx, tb + 0.4, tw, ah + 0.0005, after, "body_ink", name="after line")
        n = len(metrics)
        gap = 0.2
        th = (tb - top - (n - 1) * gap) / n
        self._metric_tiles(s, metrics, [(tx, top + i * (th + gap), tw, th) for i in range(n)], value_style="tile")
        self.card_chart(s, cx, top, cw, BOTTOM - top, chart, notes=chart_notes)
        return s

    # ------------------------------------------------------------------ demo grid
    def demo_grid(self, s, cells):
        """Fallback media on columns 5-12, in bullet order: 'band' renders run full width as
        rounded strips (they carry their own dark surround), then the last two cells sit side by
        side ('window' rounded shot, 'notch' mini MacBook lid)."""
        dx, dw = col_x(5), span(8)
        g = 24 * PX
        bands = [(k, f) for k, f in cells if k == "band"]
        rest = [(k, f) for k, f in cells if k != "band"]
        cw = (dw - g) / 2
        lid_h = self.macbook_height(cw, pad_px=8, with_base=False)
        bh = []
        for _, f in bands:
            iw, ih = self.image_size(self.image_path(f))
            bh.append(dw * ih / iw)
        tot = sum(bh) + g * len(bands) + (lid_h if rest else 0)
        if tot > BOTTOM - TOP + 1e-6:
            raise LayoutError(f"demo_grid {s.page}: media need {tot:.2f} in")
        y = TOP + (BOTTOM - TOP - tot) / 2
        for (_, f), h in zip(bands, bh):
            path = self.image_path(f)
            self.picture(s, path, dx, y, dw, h, radius=12 * PX, name=f"demo band {f}")
            y += h + g
        for i, (kind, f) in enumerate(rest):
            x = dx + i * (cw + g)
            if kind == "notch":
                self.macbook(s, x, y, cw, f, pad_px=8, with_base=False, name=f"demo device {f}")
            else:
                self.window_shot(s, f, x, y, cw, lid_h, name=f"demo window {f}")

    # ------------------------------------------------------------------ proof (tiles + window)
    def proof(self, title, *, eyebrow, lead, metrics, lines, image, notes, footer):
        """2 x 2 tiles (cols 1-6) beside the window render (cols 7-12); the three body lines
        run as three columns along the bottom."""
        s = self._new("metrics_dashboard", "mist", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1_dash", title_lines=1, lead_lines=2)
        top = y + 0.3
        need = max(self.para_h(t, "body", span(4)) for t in lines) + 0.0005
        lt = BOTTOM - need - 0.15
        for i, t in enumerate(lines):
            x = col_x(1 + 4 * i)
            self.line(s, x, lt, x + span(4), lt, name="line rule")
            self.text(s, x, lt + 0.15, span(4), need, t, "body", name="line")
        tb = lt - 0.3
        lx, lw = col_x(1), span(6)
        gap = 0.2
        tw, th = (lw - gap) / 2, (tb - top - gap) / 2
        boxes = [(lx + (i % 2) * (tw + gap), top + (i // 2) * (th + gap), tw, th) for i in range(4)]
        self._metric_tiles(s, metrics, boxes)
        path = self.image_path(image)
        f = self.fit_image(path, col_x(7), top, span(6), tb - top, mode="contain", align="r", valign="t")
        self.picture(s, path, f["x"], f["y"], f["w"], f["h"], radius=0.1, name=f"ui audit window {image}")
        return s

    # ------------------------------------------------------------------ timeline + cards
    def rail(self, s, x0, w, y, nodes, *, text_style="body", date_style="h3", name="rail"):
        """Dates above a hairline rail, dots on it, text under it. Returns the bottom y.
        node: {time, text, active, tbd}."""
        n = len(nodes)
        cw = (w - (n - 1) * GUTTER) / n
        dh = self.para_h("Ag", date_style, cw) + 0.0005
        th = max(self.para_h(nd["text"], text_style, cw) for nd in nodes) + 0.0005
        ry = y + dh + 0.16
        rail = self.line(s, x0, ry, x0 + w, ry, name=f"{name} rail")
        for i, nd in enumerate(nodes):
            xi = x0 + i * (cw + GUTTER)
            self.text(s, xi, y, cw, dh, nd["time"], date_style, name=f"{name} date")
            d = 0.16 if nd.get("active") else 0.11
            if nd.get("tbd"):
                dot = self.rect(s, xi, ry - d / 2, d, d, oval=True, fill="white", hosts=(rail,), name=f"{name} dot")
                self.dashed(s, dot, width=1.0)
            else:
                self.rect(s, xi, ry - d / 2, d, d, oval=True, fill="red" if nd.get("active") else "ink2",
                          hosts=(rail,), name=f"{name} dot")
            self.text(s, xi, ry + 0.2, cw, th, nd["text"], text_style, name=f"{name} text")
        return ry + 0.2 + th

    def timeline_cards(self, title, *, eyebrow, lead, nodes, cards, notes, footer):
        s = self._new("timeline", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        b = self.rail(s, MARGIN, CONTENT_W, y + 0.45, nodes)
        padx, pady = 28 * PX, 26 * PX
        ps = [self.P(c, "summary", weight="regular") for c in cards]
        need = max(self.measure([p], span(4) - 2 * padx)[0] for p in ps)
        ch = need + 2 * pady
        top = BOTTOM - ch
        if top < b + 0.35:
            raise LayoutError(f"timeline_cards {s.page}: cards collide with the rail")
        for i, p in enumerate(ps):
            x = col_x(1 + 4 * i)
            cd = self.card(s, x, top, span(4), ch, name="why card")
            self.text(s, x + padx, top + pady, span(4) - 2 * padx, need + 0.0005, [p], "summary", parent=cd,
                      name="why card text")
        return s

    def two_rails(self, title, *, eyebrow, lead, nodes, notes, footer):
        s = self._new("timeline", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        half = math.ceil(len(nodes) / 2)
        b = self.rail(s, MARGIN, CONTENT_W, y + 0.45, nodes[:half], name="rail 1")
        b2 = self.rail(s, MARGIN, CONTENT_W, b + 0.45, nodes[half:], name="rail 2")
        if b2 > BOTTOM + 1e-6:
            raise LayoutError(f"two_rails {s.page}: {b2:.2f} > {BOTTOM:.2f}")
        return s

    # ------------------------------------------------------------------ market
    def tiles_and_chart(self, title, *, eyebrow, lead, metrics, chart, chart_notes, notes, footer):
        s = self._new("metrics_dashboard", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1_dash", title_lines=1, lead_lines=2)
        top = y + 0.3
        lx, lw = col_x(1), span(6)
        gap = 0.2
        tw, th = (lw - gap) / 2, (BOTTOM - top - gap) / 2
        boxes = [(lx + (i % 2) * (tw + gap), top + (i // 2) * (th + gap), tw, th) for i in range(4)]
        self._metric_tiles(s, metrics, boxes)
        self.card_chart(s, col_x(7), top, span(6), BOTTOM - top, chart, notes=chart_notes)
        return s

    def funnel(self, title, *, eyebrow, lead, rows, notes, footer):
        """TAM / SAM / SOM as three cards: layer label, figure, definition, source."""
        s = self._new("metrics_dashboard", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=2)
        top = y + 0.45
        cw = span(4)
        padx, pady = 32 * PX, 28 * PX
        iw = cw - 2 * padx
        lab_h = TYPE["label"].size * TYPE["label"].lh / 72 + 0.0005
        fh = max(self.para_h(r[2], "tile", iw) for r in rows) + 0.0005
        dh = max(self.para_h(r[1], "body_ink", iw) for r in rows) + 0.0005
        sh = max(self.para_h("Source: " + r[3], "small", iw) for r in rows) + 0.0005
        ch = pady + lab_h + 0.16 + fh + 0.2 + dh + 0.3 + sh + pady
        if top + ch > BOTTOM + 1e-6:
            raise LayoutError(f"funnel {s.page}: {top + ch:.2f} > {BOTTOM}")
        for i, r in enumerate(rows):
            x = col_x(1 + 4 * i)
            cd = self.card(s, x, top, cw, ch, name="funnel card")
            yy = top + pady
            self.text(s, x + padx, yy, iw, lab_h, r[0], "label", parent=cd, name="funnel layer")
            yy += lab_h + 0.16
            self.text(s, x + padx, yy, iw, fh, r[2], "tile", parent=cd, name="funnel figure")
            yy += fh + 0.2
            self.text(s, x + padx, yy, iw, dh, r[1], "body_ink", parent=cd, name="funnel definition")
            yy += dh + 0.3
            self.line(s, x + padx, yy - 0.15, x + cw - padx, yy - 0.15, parent=cd, name="funnel rule")
            self.text(s, x + padx, yy, iw, sh, "Source: " + r[3], "small", parent=cd, name="funnel source")
        return s

    def table_slide(self, title, *, eyebrow, lead, columns, rows, widths, notes, footer, after=(),
                    cell="cell", first="cell_ink", bg="white"):
        s = self._new("comparison_table", bg, notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=2)
        ah = sum(self.para_h(a, "small", CONTENT_W) for a in after) + (0.3 if after else 0)
        b = self._table(s, MARGIN, y + 0.35, CONTENT_W, columns, rows, widths=widths,
                        max_bottom=BOTTOM - ah, cell=cell, first=first)
        if after:
            self.text(s, MARGIN, b + 0.3, CONTENT_W, ah - 0.3 + 0.0005, list(after), "small", name="table note")
        return s

    # ------------------------------------------------------------------ competition matrix
    def matrix(self, title, *, eyebrow, lead, columns, rows, highlight, legend, notes, footer):
        s = self._new("comparison_table", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        n = len(columns)
        first_w, cgap = 2.15, 0.12
        pw = (CONTENT_W - first_w - (n - 1) * cgap) / (n - 1)
        xs = [MARGIN] + [MARGIN + first_w + cgap + i * (pw + cgap) for i in range(n - 1)]
        ws = [first_w] + [pw] * (n - 1)
        hi = columns.index(highlight)
        top = y + 0.35
        bar_h = 4 * PX
        self.rect(s, xs[hi] - cgap / 2, top, ws[hi] + cgap, bar_h, fill="red", name="highlight bar")
        head_y = top + bar_h + 0.1
        hh = max(self.para_h(c, "cell_sm_ink", pw) for c in columns[1:]) + 0.0005
        for j, c in enumerate(columns):
            if j:
                self.text(s, xs[j], head_y, ws[j], hh, c, "cell_sm_ink", align="c", name="matrix head")
        yy = head_y + hh + 0.1
        self.line(s, MARGIN, yy, MARGIN + CONTENT_W, yy, color="ink", name="matrix head rule")
        legend_h = self.para_h(legend, "note", CONTENT_W) + 0.0005
        vpad = 0.08
        for ri, r in enumerate(rows):
            price = r[0] == "Price"
            st = "cell_sm" if price else "cell"
            rh = 0.42
            if price:
                rh = max(rh, max(self.para_h(v, st, ws[j]) for j, v in enumerate(r) if j and v) + 2 * vpad)
            rh = max(rh, self.para_h(r[0], "cell_ink", first_w) + 2 * vpad)
            if price:
                yy += 0.06
            self.text(s, xs[0], yy + vpad, first_w, rh - 2 * vpad, r[0], "cell_ink", name=f"matrix r{ri + 1} job")
            for j, v in enumerate(r[1:], start=1):
                if not v:
                    continue
                if price:
                    self.text(s, xs[j], yy + vpad, ws[j], rh - 2 * vpad, v, "cell_sm", align="c",
                              color="ink" if j == hi else None, name=f"matrix r{ri + 1}c{j + 1}")
                else:
                    gh = TYPE["body_ink"].size * TYPE["body_ink"].lh / 72 + 0.0005
                    self.text(s, xs[j], yy + (rh - gh) / 2, ws[j], gh, v, "body_ink", align="c",
                              name=f"matrix r{ri + 1}c{j + 1}")
            yy += rh
            if ri != len(rows) - 1:
                self.line(s, MARGIN, yy, MARGIN + CONTENT_W, yy,
                          color="hairline", name="matrix rule")
        if yy + 0.25 + legend_h > BOTTOM + 1e-6:
            raise LayoutError(f"matrix {s.page}: {yy + 0.25 + legend_h:.2f} > {BOTTOM}")
        self.text(s, MARGIN, yy + 0.25, CONTENT_W, legend_h, legend, "note", name="matrix legend")
        return s

    def positioning(self, title, *, eyebrow, lead, chart, lines, notes, footer):
        s = self._new("comparison_table", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        top = y + 0.35
        cap = chart["label"][0].upper() + chart["label"][1:] + "."
        cap_h = self.para_h(cap, "note", span(6)) + 0.0005
        title_h = self.para_h(chart["title"], "small_ink", span(6)) + 0.0005
        size = BOTTOM - top - cap_h - title_h - 0.3
        mw = span(6)
        mx = col_x(1)
        self.text(s, mx, top, span(6), title_h, chart["title"], "small_ink", name="map title")
        my = top + title_h + 0.15
        box = self.card(s, mx, my, mw, size, name="positioning map")
        self.text(s, mx, my + size + 0.15, span(6), cap_h, cap, "note", name="map caption")
        xa, xb = [t.strip() for t in chart["x_axis"].split("→")]
        ya, yb = [t.strip() for t in chart["y_axis"].split("→")]
        lab = TYPE["label"]
        lh = lab.size * lab.lh / 72 + 0.0005
        m = 0.12
        cx, cy = mx + mw / 2, my + size / 2
        # axis end labels, then axis lines that stop short of them
        wy_b = self.text_width(yb, "label") + 0.04
        wy_a = self.text_width(ya, "label") + 0.04
        wx_a = self.text_width(xa, "label") + 0.04
        wx_b = self.text_width(xb, "label") + 0.04
        self.text(s, cx - wy_b / 2, my + m, wy_b, lh, yb, "label", align="c", parent=box, name="axis label")
        self.text(s, cx - wy_a / 2, my + size - m - lh, wy_a, lh, ya, "label", align="c", parent=box, name="axis label")
        self.text(s, mx + m, cy - lh / 2, wx_a, lh, xa, "label", parent=box, name="axis label")
        self.text(s, mx + mw - m - wx_b, cy - lh / 2, wx_b, lh, xb, "label", align="r", parent=box, name="axis label")
        vl = self.line(s, cx, my + m + lh + 0.08, cx, my + size - m - lh - 0.08, parent=box, name="axis y")
        hl = self.line(s, mx + m + wx_a + 0.08, cy, mx + mw - m - wx_b - 0.08, cy, parent=box, hosts=(vl,),
                       name="axis x")
        pad = 0.45
        inner = size - 2 * pad
        padx_ = 1.0
        inner_x = mw - 2 * padx_
        side = {"Wispr Flow": "l", "SparkV": "b"}
        for p in chart["points"]:
            px_, py_ = mx + padx_ + p["x"] * inner_x, my + pad + (1 - p["y"]) * inner
            d = 0.16 if p.get("accent") else 0.1
            dot = self.rect(s, px_ - d / 2, py_ - d / 2, d, d, oval=True, fill="red" if p.get("accent") else "ink2",
                            parent=box, hosts=(vl, hl), name=f"map dot {p['label']}")
            text = p["label"] + (f" ({p['note']})" if p.get("note") else "")
            st = "small_ink" if p.get("accent") else "small"
            tw = self.text_width(text, st) + 0.04
            th = TYPE[st].size * TYPE[st].lh / 72 + 0.0005
            sd_ = side.get(p["label"], "r")
            ty = py_ - th / 2
            if sd_ == "l":
                tx = px_ - d / 2 - 0.06 - tw
            elif sd_ == "b":
                tx = min(px_ - tw / 2, mx + mw - 0.08 - tw)
                ty = py_ + d / 2 + 0.04
            else:
                tx = px_ + d / 2 + 0.06
            self.text(s, tx, ty, tw, th, text, st, align={"l": "r", "b": "c"}.get(sd_, "l"),
                      fill="mist", parent=box, hosts=(vl, hl, dot), name=f"map label {p['label']}")
        rx, rw = col_x(8), span(5)
        self.rule_rows(s, rx, my, rw, lines, "lead", gap=18 * PX)
        return s

    # ------------------------------------------------------------------ unit economics
    def unit_economics(self, title, *, eyebrow, lead, chart, chart_notes, metrics, lines, notes, footer):
        s = self._new("metrics_dashboard", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1_dash", title_lines=1, lead_lines=2)
        top = y + 0.3
        rx, rw = col_x(9), span(4)
        hs = [self.para_h(t, "body", rw) for t in lines]
        rows_h = sum(hs) + (2 * len(lines) - 1) * 12 * PX + 0.002 * len(lines)
        tb = BOTTOM - rows_h - 0.15
        gap = 0.2
        th = (tb - top - gap) / 2
        self._metric_tiles(s, metrics, [(rx, top + i * (th + gap), rw, th) for i in range(2)])
        self.rule_rows(s, rx, tb + 0.15, rw, lines, "body", gap=12 * PX)
        wf_in = chart
        vals = wf_in["series"][0]["values"]
        tot = set(wf_in["total_indices"])
        base, step, run = [], [], 0.0
        for i, v in enumerate(vals):
            if i in tot:
                base.append(0.0)
                step.append(v)
                run = v
            else:
                run += v
                base.append(run)
                step.append(-v)
        axis_max = sd._nice_max(max(vals))
        capv = axis_max * 0.09
        cats = wf_in["categories"]
        last = len(cats) - 1
        colors = ["ink" if i == 0 else "red" if i == last else "bar" for i in range(len(cats))]
        labels = []
        for i, v in enumerate(vals):
            t = fmt_num(abs(v), wf_in["number_format"])
            labels.append(("−" + t) if v < 0 else t)
        zero_i = [i for i, v in enumerate(vals) if v == 0]
        label_colors = ["red" if (i in zero_i or i == last) else "ink" for i in range(len(cats))]
        spec = {"type": "column_stacked", "title": wf_in["title"], "categories": cats,
                "number_format": wf_in["number_format"], "axis_max": axis_max, "legend": False,
                "series": [{"name": "base", "values": base, "color": "mist"},
                           {"name": "USD", "values": step, "color": "bar",
                            "point_colors": {i: c for i, c in enumerate(colors)}},
                           {"name": "label", "values": [capv] * len(cats), "color": "mist"}],
                "source": wf_in["source_line"]}
        wf = {"base": base, "step": step, "cap": capv, "colors": colors, "labels": labels,
              "label_colors": label_colors}
        self.card_chart(s, col_x(1), top, span(8), BOTTOM - top, spec, notes=chart_notes, waterfall=wf)
        return s

    def per_plan(self, title, *, eyebrow, lead, table, chart, notes, footer):
        s = self._new("metrics_dashboard", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1_dash", title_lines=1, lead_lines=2)
        top = y + 0.3
        lx, lw = col_x(1), span(7)
        th = self.para_h(table["title"], "small_ink", lw) + 0.0005
        self.text(s, lx, top, lw, th, table["title"], "small_ink", name="table title")
        basis = "Basis: " + table["basis"]
        bh = self.para_h(basis, "small", lw) + 0.0005
        b = self._table(s, lx, top + th + 0.2, lw, table["columns"], table["rows"], widths=[2.6, 1.0, 1.3, 1.5],
                        max_bottom=BOTTOM - bh - 0.25)
        self.text(s, lx, b + 0.25, lw, bh, basis, "small", name="table basis")
        self.card_chart(s, col_x(8), top, span(5), BOTTOM - top, chart["spec"], notes=chart["notes"],
                        title_color="ink3", dashed=True)
        return s

    # ------------------------------------------------------------------ go to market, channels
    def gtm_channels(self, title, *, eyebrow, lead, points, channels, targets, notes, footer):
        s = self._new("timeline", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=1, lead_lines=1)
        yy = y + 0.45
        need = max(self.para_h(p, "body_ink", span(4)) for p in points) + 0.0005
        for i, p in enumerate(points):
            x = col_x(1 + 4 * i)
            self.line(s, x, yy, x + span(4), yy, name="point rule")
            self.text(s, x, yy + 0.15, span(4), need, p, "body_ink", name="point")
        yy += 0.15 + need + 0.4
        lab_h = TYPE["label"].size * TYPE["label"].lh / 72 + 0.0005
        self.text(s, MARGIN, yy, CONTENT_W, lab_h, "Channels", "label", name="row label")
        yy += lab_h + 0.14
        ch_h = TYPE["small_ink"].size * TYPE["small_ink"].lh / 72 + 2 * 9 * PX
        x = MARGIN
        for c in channels:
            w = self.chip_width(c)
            if x + w > MARGIN + CONTENT_W + 1e-6:
                x = MARGIN
                yy += ch_h + 0.14
            self.chip(s, x, yy, c, line="ink", name="channel chip")
            x += w + 0.14
        yy += ch_h + 0.4
        self.text(s, MARGIN, yy, CONTENT_W, lab_h, "Targets", "label", name="row label")
        yy += lab_h + 0.14
        n = len(targets)
        tw = (CONTENT_W - (n - 1) * GUTTER) / n
        th = max(self.para_h(t, "body", tw - 0.3) for t in targets) + 0.36
        if yy + th > BOTTOM + 1e-6:
            raise LayoutError(f"gtm_channels {s.page}: {yy + th:.2f} > {BOTTOM}")
        for i, t in enumerate(targets):
            x = MARGIN + i * (tw + GUTTER)
            box = self.rect(s, x, yy, tw, th, fill="white", radius=14 * PX, name="target tile")
            self.dashed(s, box)
            hh = self.para_h(t, "body", tw - 0.3) + 0.0005
            self.text(s, x + 0.15, yy + (th - hh) / 2, tw - 0.3, hh, t, "body", align="c", parent=box,
                      name="target text")
        return s

    # ------------------------------------------------------------------ traction
    def traction(self, title, *, eyebrow, lead, metrics, tbd, lines, notes, footer):
        s = self._new("metrics_dashboard", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1_dash", title_lines=1, lead_lines=1)
        lab_h = TYPE["label"].size * TYPE["label"].lh / 72 + 0.0005
        yy = y + 0.3
        self.text(s, MARGIN, yy, CONTENT_W, lab_h, "Verified", "label", name="row label")
        yy += lab_h + 0.12
        n = len(metrics)
        tw = span(3)
        padx, pady = 28 * PX, 22 * PX
        mk = 0.12
        lw = tw - 2 * padx - mk - 0.1
        lab_need = max(self.para_h(m["label"], "small", lw) for m in metrics) + 0.0005
        ts = TYPE["kpi"]
        vh = ts.size * 0.8 * ts.lh / 72 + 0.0005
        th = pady + vh + 0.08 + lab_need + pady
        for i, m in enumerate(metrics):
            x = col_x(1 + 3 * i)
            cd = self.card(s, x, yy, tw, th, name="metric card")
            p = self.P(m["value"], "kpi").scaled(0.8)
            self.text(s, x + padx, yy + pady, tw - 2 * padx, vh, [p], "kpi", parent=cd, name="metric value")
            ly = yy + pady + vh + 0.08
            self.mark(s, "check", x + padx, ly + 0.025, mk, color=OK_GREEN, width=1.4, parent=cd, name="ok check")
            self.text(s, x + padx + mk + 0.1, ly, lw, lab_need, m["label"], "small", parent=cd, name="metric label")
        yy += th + 0.35
        self.text(s, MARGIN, yy, CONTENT_W, lab_h, "What we will report", "label", name="row label")
        yy += lab_h + 0.12
        k = len(tbd)
        bw = (CONTENT_W - (k - 1) * GUTTER) / k
        bh = max(self.para_h(t, "body", bw - 0.3) for t in tbd) + 0.32
        for i, t in enumerate(tbd):
            x = MARGIN + i * (bw + GUTTER)
            box = self.rect(s, x, yy, bw, bh, fill="white", radius=14 * PX, name="tbd tile")
            self.dashed(s, box)
            hh = self.para_h(t, "body", bw - 0.3) + 0.0005
            self.text(s, x + 0.15, yy + (bh - hh) / 2, bw - 0.3, hh, t, "body", align="c", parent=box, name="tbd text")
        yy += bh + 0.35
        need = max(self.para_h(t, "body", span(4)) for t in lines) + 0.0005
        if yy + 0.15 + need > BOTTOM + 1e-6:
            raise LayoutError(f"traction {s.page}: {yy + 0.15 + need:.2f} > {BOTTOM}")
        top = BOTTOM - 0.15 - need
        for i, t in enumerate(lines):
            x = col_x(1 + 4 * i)
            self.line(s, x, top, x + span(4), top, name="line rule")
            self.text(s, x, top + 0.15, span(4), need, t, "body_ink" if i < 2 else "body", name="line")
        return s

    # ------------------------------------------------------------------ roadmap lane
    def lane(self, title, *, eyebrow, lead, label, caption, items, dashed, notes, footer):
        s = self._new("timeline", "white", notes, footer, None, False)
        y = self._header(s, eyebrow, title, lead, h1="h1", title_lines=2, lead_lines=1)
        n = len(items)
        cw = (CONTENT_W - (n - 1) * GUTTER) / n
        padx, pady = 26 * PX, 24 * PX
        iw = cw - 2 * padx
        lab_h = TYPE["label"].size * TYPE["label"].lh / 72 + 0.0005
        t_h = max(self.para_h(it["title"], "h3", iw) for it in items) + 0.0005
        d_h = max(self.para_h(it["detail"], "body", iw) for it in items) + 0.0005
        ch = pady + lab_h + 0.14 + t_h + 0.1 + d_h + pady
        cap_h = self.para_h(caption, "body_ink", CONTENT_W) + 0.0005
        block = lab_h + 0.12 + cap_h + 0.3 + ch
        top = y + 0.45
        if top + block > BOTTOM + 1e-6:
            raise LayoutError(f"lane {s.page}: {top + block:.2f} > {BOTTOM}")
        top += (BOTTOM - top - block) / 2
        self.text(s, MARGIN, top, CONTENT_W, lab_h, label, "label", name="lane label")
        self.text(s, MARGIN, top + lab_h + 0.12, CONTENT_W, cap_h, caption, "body_ink", name="lane caption")
        cy = top + lab_h + 0.12 + cap_h + 0.3
        for i, it in enumerate(items):
            x = MARGIN + i * (cw + GUTTER)
            if dashed:
                cd = self.rect(s, x, cy, cw, ch, fill="white", radius=22 * PX, name="lane card")
                self.dashed(s, cd)
            else:
                cd = self.card(s, x, cy, cw, ch, name="lane card")
            yy = cy + pady
            self.text(s, x + padx, yy, iw, lab_h, f"{i + 1:02d}", "label", parent=cd, name="lane index")
            yy += lab_h + 0.14
            self.text(s, x + padx, yy, iw, t_h, it["title"], "h3", parent=cd, name="lane title")
            yy += t_h + 0.1
            self.text(s, x + padx, yy, iw, d_h, it["detail"], "body", parent=cd, name="lane detail")
        return s

    # ------------------------------------------------------------------ team
    def team(self, title, *, lead, cards, studio, notes, footer):
        s = self._new("minimal_body", "white", notes, footer, None, False)
        y = self._header(s, None, title, lead, h1="h1_large", title_lines=1, lead_lines=1)
        st_h = self.para_h(studio, "body", CONTENT_W) + 0.0005
        top = y + 0.4
        bot = BOTTOM - st_h - 0.3
        cw = span(4)
        padx, pady = 28 * PX, 28 * PX
        iw = cw - 2 * padx
        fh = [max(self.para_h(c[k], st, iw) for c in cards) + 0.0005
              for k, st in (("name", "h3"), ("role", "body"), ("line", "body"))]
        for i, c in enumerate(cards):
            x = col_x(1 + 4 * i)
            cd = self.card(s, x, top, cw, bot - top, name="team card")
            ph = min(1.35, (bot - top) * 0.42)
            px_ = x + (cw - ph) / 2
            circ = self.rect(s, px_, top + pady, ph, ph, oval=True, fill="white", parent=cd, name="photo slot")
            self.dashed(s, circ)
            lh = self.para_h(c["photo"], "small", ph * 0.7) + 0.0005
            self.text(s, px_ + ph * 0.15, top + pady + (ph - lh) / 2, ph * 0.7, lh, c["photo"], "small", align="c",
                      parent=circ, name="photo placeholder")
            yy = top + pady + ph + 0.25
            for txt, st, hh in ((c["name"], "h3", fh[0]), (c["role"], "body", fh[1]), (c["line"], "body", fh[2])):
                self.text(s, x + padx, yy, iw, hh, txt, st, align="c", parent=cd, name="team field")
                yy += hh + 0.08
            if yy - 0.08 > bot - pady + 1e-6:
                raise LayoutError(f"team {s.page}: card fields overflow")
        self.text(s, MARGIN, BOTTOM - st_h, CONTENT_W, st_h, studio, "body", align="c", name="studio line")
        return s

    # ------------------------------------------------------------------ the ask
    def ask(self, title, *, lead, lines, button, beside, chart, legend, notes, footer):
        s = self._new("quote_or_cta", "mist", notes, footer, None, False)
        lx, lw = col_x(1), span(7)
        ts = TYPE["h1_large"]
        y = TOP + 0.35
        t = self.text(s, lx, y, lw, None, title, "h1_large", max_h=2 * ts.size * ts.lh / 72 + 0.001, name="headline")
        l_ = self.text(s, lx, t.b + 26 * PX, lw, None, lead, "lead", max_h=3 * TYPE["lead"].size * TYPE["lead"].lh / 72 + 0.001,
                       name="lead")
        b = self.rule_rows(s, lx, l_.b + 0.45, lw, lines, "body_ink", gap=16 * PX)
        bh = TYPE["button"].size * TYPE["button"].lh / 72 + 2 * 15 * PX
        btn = self.button(s, lx, b + 0.35, button, name="cta")
        bw_ = self.text_width(beside, "small") + 0.05
        sh = TYPE["small"].size * TYPE["small"].lh / 72 + 0.0005
        self.text(s, btn.r + 28 * PX, btn.y + (bh - sh) / 2, bw_, sh, beside, "small", name="cta beside")
        self.card_chart(s, col_x(9), TOP + 0.35, span(4), BOTTOM - TOP - 0.35, chart, legend_rows=legend,
                        fill="white")
        return s

    # ------------------------------------------------------------------ appendix
    def feature_grid(self, title, *, lead, files, captions, notes, footer, page):
        s = self._new("feature_showcase", "mist", notes, footer, page, True)
        y = self._header(s, None, title, lead, h1="h1_dash", title_lines=1, lead_lines=1)
        top = y + 0.3
        cols_, rows_ = 4, math.ceil(len(files) / 4)
        cw = span(3)
        cap_h = TYPE["small_ink"].size * TYPE["small_ink"].lh / 72 + 0.0005
        hs = []
        for f in files:
            iw, ih = self.image_size(self.image_path(f))
            hs.append(cw * ih / iw)
        row_h = [max(hs[r * cols_:(r + 1) * cols_]) for r in range(rows_)]
        rg = 0.3
        tot = sum(row_h) + rows_ * (cap_h + 0.1) + (rows_ - 1) * rg
        if top + tot > BOTTOM + 1e-6:
            raise LayoutError(f"feature_grid {s.page}: {top + tot:.2f} > {BOTTOM}")
        yy = top + (BOTTOM - top - tot) / 2
        for r in range(rows_):
            for k in range(cols_):
                i = r * cols_ + k
                if i >= len(files):
                    break
                x = col_x(1 + 3 * k)
                path = self.image_path(files[i])
                self.picture(s, path, x, yy, cw, hs[i], radius=10 * PX, name=f"grid render {files[i]}")
                self.text(s, x, yy + row_h[r] + 0.1, cw, cap_h, captions[i], "small_ink", align="c",
                          name="grid caption")
            yy += row_h[r] + cap_h + 0.1 + rg
        return s

    def sources(self, title, *, lead, sources, notes, footer, first_page=2):
        """Two-column source list; runs over as many hidden slides as it needs."""
        cw = span(6)

        def break_url(u):
            """Hard line breaks after '/' or '-' so a long URL never runs past the column."""
            parts = re.split(r"(?<=[/-])", u)
            lines, cur = [], ""
            for p_ in parts:
                if cur and self.text_width(cur + p_, "note") > cw - 0.05:
                    lines.append(cur)
                    cur = p_
                else:
                    cur += p_
            lines.append(cur)
            return "\n".join(lines)
        entries = []
        for so in sources:
            entries.append([self.P(so["title"], "small_ink"), self.P(so["note"], "small"),
                            self.P(break_url(so["url"]), "note")])
        gap = 0.16
        pages, cur, col_i, used = [], [[], []], 0, 0.0
        avail = None
        made = []
        i = 0
        page_no = first_page
        while i < len(entries):
            s = self._new("appendix_table", "white", notes, footer, f"A{page_no}", True)
            y = self._header(s, None, title, lead if page_no == first_page else None, h1="h1_dash",
                             title_lines=1, lead_lines=1)
            top = y + 0.35
            for c in range(2):
                yy = top
                x = col_x(1 + 6 * c)
                while i < len(entries):
                    ps = entries[i]
                    h = self.measure(ps, cw)[0] + 0.0005
                    if yy + h > BOTTOM + 1e-6:
                        break
                    self.text(s, x, yy, cw, h, ps, "small", name="source entry")
                    yy += h + gap
                    i += 1
            made.append(s)
            page_no += 1
            if len(made) > 6:
                raise LayoutError("sources: runaway pagination")
        return made


# =========================================================================== build
def build() -> InvestorDeck:
    o = load_outline("investor")
    S = {s["id"]: s for s in o["slides"]}
    SRC = {x["id"]: x for x in o["sources"]}
    d = InvestorDeck(o["deck"]["title"], footer_default="")

    # 1 Title ---------------------------------------------------------------------------
    s = S["title"]
    d.investor_title(H(s), eyebrow=s["body"][0], lead=s["subtitle"], line=" · ".join(s["body"][1:]),
                     image=s["images"][0]["file"], notes=s["notes"])

    # 2 Problem: cost chart C1 + three tiles ---------------------------------------------
    s = S["problem"]
    c = s["chart"]
    idx = {k: c["categories"].index(k) for k in c["highlight"]}
    pc = {idx[k]: ("red" if v == "accent" else "ink") for k, v in c["highlight"].items()}
    d.chart_and_tiles(H(s), eyebrow="Problem", lead=s["subtitle"],
                      chart={"type": "bar", "title": c["title"], "categories": c["categories"],
                             "number_format": c["number_format"],
                             "series": [{"name": c["series"][0]["name"], "values": c["series"][0]["values"],
                                         "color": "bar", "point_colors": pc}],
                             "source": c["source_line"]},
                      metrics=[{"value": t["value"], "label": t["label"], "accent": t.get("accent", False)}
                               for t in s["kpi_tiles"]],
                      after=s["body"][2], notes=s["notes"],
                      footer="Year one: subscriptions for 12 months, one-time apps at their list price.")

    # 3 Solution -------------------------------------------------------------------------
    s = S["solution"]
    bl = [re.sub(r"^([^:]+):", r"**\1:**", b) for b in s["body"]]
    d.feature_showcase(H(s), eyebrow="Solution", lead=s["subtitle"], bullets=bl,
                       image=s["images"][0]["file"], frame="notch", notes=s["notes"],
                       footer="Real render of the current build, sample data.")

    # 4 Live demo: text + 2 x 2 fallback grid --------------------------------------------
    s = S["live_demo"]
    sl = d.feature_showcase(H(s), eyebrow="Live demo", lead=s["subtitle"], bullets=s["body"], image=None,
                            notes=s["notes"],
                            footer="Fallback stills for the live demo. Real renders of the current build, sample data.")
    cells = []
    for im in s["images"]:
        f, kind = SAFE_IMAGE.get(im["file"], im["file"]), im["mockup"]
        if im["file"] == "16-universal-search-panel.png":
            kind = "notch"          # the substitute is a drawer render
        cells.append((kind, f))
    d.demo_grid(sl, cells)

    # 5 Meetings --------------------------------------------------------------------------
    s = S["meetings"]
    d.feature_showcase(H(s), eyebrow="Meetings", lead=s["subtitle"], bullets=s["body"],
                       image=SAFE_IMAGE.get(s["images"][0]["file"], s["images"][0]["file"]), frame="window",
                       notes=s["notes"],
                       footer="Rendered from a saved meeting with sample data. Live audio capture: pre-launch live check.")

    # 6 Proof -------------------------------------------------------------------------------
    s = S["proof_ui_audit"]
    d.proof(H(s), eyebrow="Proof", lead=s["subtitle"],
            metrics=[{"value": t["value"], "label": t["label"], "accent": t.get("accent", False)}
                     for t in s["kpi_tiles"]],
            lines=s["body"], image=s["images"][0]["file"], notes=s["notes"], footer=s["source_line"])

    # 7 Why now -----------------------------------------------------------------------------
    s = S["why_now"]
    nodes = []
    for t in s["timeline"]:
        nodes.append({"time": t["when"], "text": t["what"], "active": t["when"] == "Jun 2025",
                      "tbd": t["when"] == "Next"})
    cards = [s["body"][0].replace("SparkV pays $0.", "SparkV pays ==$0==."), s["body"][1], s["body"][2]]
    d.timeline_cards(H(s), eyebrow="Why now", lead=s["subtitle"], nodes=nodes, cards=cards, notes=s["notes"],
                     footer="Sources: Macworld (Oct 2021); MobileSyrup (Jun 2023); TechCrunch (Jun 2025); "
                            "IDC preliminary via 9to5Mac (Jan 2026); SparkV business plan.")

    # 8a Market: tiles + chart C2 --------------------------------------------------------------
    s = S["market"]
    c = s["chart"]
    n1, n2 = split_notes(s["notes"], 3)
    cats = [cat + " (derived)" if cat == "2024" else cat for cat in c["categories"]]
    d.tiles_and_chart(H(s), eyebrow="Market", lead=s["subtitle"],
                      metrics=[{"value": t["value"], "label": t["label"], "accent": t.get("accent", False)}
                               for t in s["kpi_tiles"]],
                      chart={"type": "column", "title": c["title"], "categories": cats,
                             "number_format": c["number_format"],
                             "series": [{"name": c["series"][0]["name"], "values": c["series"][0]["values"],
                                         "color": "bar", "point_colors": {1: "red"}}],
                             "source": c["source_line"]},
                      chart_notes=c["annotations"], notes=n1,
                      footer="Setapp 2023 surveyed Setapp's own audience (n=1,241, mostly US), not all Mac owners.")

    # 8b Market: sizing funnel T8 ----------------------------------------------------------------
    t = s["table"]

    def src_name(v):
        return SRC[v]["title"] if v in SRC else v
    rows = [[r[0], r[1], r[2], src_name(r[3])] for r in t["rows"]]
    d.funnel(H(s), eyebrow="Market · " + t["title"], lead=None, rows=rows, notes=n2,
                  footer="Figures and definitions from the sizing funnel; [TBD] fields stay open until sourced.")

    # 9a Competition: matrix T9 ------------------------------------------------------------------
    s = S["competition"]
    n1, n2 = split_notes(s["notes"], 4)
    t = s["table"]
    d.matrix(H(s), eyebrow="Competition", lead=s["subtitle"], columns=t["columns"], rows=t["rows"],
             highlight=t["highlight_column"], legend=t["legend"], notes=n1,
             footer="Prices from third-party listings cited in the SparkV business plan (section 3), Oct 2026. "
                    "Re-check vendor pages before sending.")

    # 9b Competition: positioning P9 -------------------------------------------------------------
    lines = [re.sub(r"^([^:]+):", r"**\1:**", b) for b in s["body"]]
    d.positioning(H(s), eyebrow="Competition · positioning", lead=s["subtitle"], chart=s["chart"], lines=lines,
                  notes=n2, footer="Qualitative map. Price models from the listings on the previous slide.")

    # 10 Business model ------------------------------------------------------------------------------
    s = S["business_model"]
    cards = {c_["plan"]: c_ for c_ in s["table"]["cards"]}
    fr, pro, team = cards["Free"], cards["Pro"], cards["Team"]
    unit_main, unit_rest = team["unit"].split(", ", 1)
    plans = [
        {"name": "Free", "price": fr["price"], "inr": fr["price_inr"], "note": s["body"][0], "features": fr["lines"]},
        {"name": "Pro", "price": pro["price"], "unit": pro["unit"], "inr": pro["price_inr"], "note": s["body"][1],
         "features": pro["lines"], "highlight": True, "tag": pro["tag"]},
        {"name": "Team", "price": team["price"], "unit": unit_main, "inr": team["price_inr"], "inr_aside": unit_rest,
         "note": s["body"][2], "features": team["lines"]},
    ]
    d.pricing_table(H(s), plans, eyebrow="Business model", lead=s["subtitle"], aside=s["table"]["footer"],
                    notes=s["notes"], footer="Source: SparkV business plan, 22 Oct 2026. INR prices as listed there.")

    # 11a Unit economics: waterfall C3 + tiles ----------------------------------------------------------
    s = S["unit_economics"]
    n1, n2 = split_notes(s["notes"], 4)
    d.unit_economics(H(s), eyebrow="Unit economics", lead=s["subtitle"], chart=s["chart"],
                     chart_notes=s["chart"]["annotations"],
                     metrics=[{"value": t_["value"], "label": t_["label"], "accent": t_.get("accent", False)}
                              for t_ in s["kpi_tiles"]],
                     lines=s["body"][1:], notes=n1,
                     footer="Lemon Squeezy 5% + $0.50 per order. On-device AI; Groq optional with the buyer's own key.")

    # 11b Unit economics: per plan T11 + illustrative C4 -------------------------------------------------
    c2 = s["chart_2"]
    d.per_plan(H(s), eyebrow="Unit economics · per plan", lead=None, table=s["table"],
               chart={"spec": {"type": "column", "title": c2["title"], "categories": c2["categories"],
                               "number_format": c2["number_format"],
                               "series": [{"name": c2["series"][0]["name"], "values": c2["series"][0]["values"],
                                           "color": "bar"}],
                               "source": "IDC preliminary 2025 shipments (via 9to5Mac) and the SparkV business plan."},
                      "notes": c2["annotations"]},
               notes=n2, footer="Illustrative arithmetic, not a forecast.")

    # 12a Go to market: 30-day plan on two rails --------------------------------------------------------
    s = S["go_to_market"]
    n1, n2 = split_notes(s["notes"], 3)
    nodes = [{"time": t_["when"], "text": t_["what"], "active": t_["when"] == "Day 27"} for t_ in s["timeline"]]
    d.two_rails(H(s), eyebrow="Go to market", lead=s["subtitle"], nodes=nodes, notes=n1,
                footer="Source: SparkV business plan, 30-day launch plan.")

    # 12b Go to market: channels and targets --------------------------------------------------------------
    d.gtm_channels(H(s), eyebrow="Go to market · channels", lead=None, points=s["body"], channels=s["channels"],
                   targets=s["kpi_placeholders"], notes=n2, footer="Source: SparkV business plan.")

    # 13 Traction --------------------------------------------------------------------------------------------
    s = S["traction"]
    d.traction(H(s), eyebrow="Traction", lead=s["subtitle"], metrics=s["kpi_tiles"], tbd=s["tbd_tiles"],
               lines=s["body"], notes=s["notes"], footer=s["source_line"])

    # 14a / 14b Roadmap lanes -----------------------------------------------------------------------------
    s = S["roadmap"]
    n1, n2 = split_notes(s["notes"], 4)
    l1, l2 = s["lanes"]
    d.lane(H(s), eyebrow="Roadmap and risks", lead=s["subtitle"], label=l1["name"], caption=s["body"][0],
           items=l1["items"], dashed=False, notes=n1, footer="Source: SparkV business plan, launch blockers.")
    d.lane(l2["name"], eyebrow="Roadmap and risks", lead=None, label="After launch", caption=s["body"][1],
           items=l2["items"], dashed=True, notes=n2, footer="Source: SparkV business plan.")

    # 15 Team ---------------------------------------------------------------------------------------------------
    s = S["team"]
    d.team(s["header"], lead=s["subtitle"], cards=s["cards"], studio=s["body"][3], notes=s["notes"],
           footer="Studio: sparkv.si, read 22 Oct 2026.")

    # 16 The ask ---------------------------------------------------------------------------------------------------
    s = S["the_ask"]
    c = s["chart"]
    greys = ["ink3", "bar", "hairline", "track"]
    pct = {r["bucket"]: r["percent"] for r in c["data_table"]}
    url, email = [p.strip() for p in s["body"][2].split("·", 1)]
    d.ask(s["header"], lead=s["subtitle"], lines=s["body"][:2], button=url, beside=email,
          chart={"type": "doughnut", "title": c["title"], "categories": c["categories"],
                 "series": [{"name": c["series"][0]["name"], "values": [1] * len(c["categories"]),
                             "point_colors": {i: g for i, g in enumerate(greys)}}],
                 "tbd": ["[TBD: % of the round per bucket. Equal slices are placeholders, not a split.]"]},
          legend=[(cat, pct[cat], greys[i]) for i, cat in enumerate(c["categories"])],
          notes=s["notes"], footer="")

    # A1 Feature grid (hidden) ------------------------------------------------------------------------------------
    s = S["appendix_features"]
    names = {"02-widgets.png": "Widgets", "04-tray.png": "Tray", "06-today.png": "Today",
             "07-converter.png": "Converter", "08-spaces.png": "Spaces", "11-flow.png": "Flow",
             "13-clipboard.png": "Clipboard", "15-universal-search.png": "Universal Search",
             "28-focus.png": "Focus", "31-downloads.png": "Downloads", "17-reply-drafter.png": "Reply Drafter",
             "23-job-agent.png": "Job Agent"}
    files = s["grid"]["files"]
    d.feature_grid(s["header"], lead=s["subtitle"], files=files, captions=[names[f] for f in files],
                   notes=s["notes"], footer="Captions are the render file names.", page="A1")

    # A2+ Sources (hidden) ------------------------------------------------------------------------------------------
    s = S["appendix_sources"]
    d.sources(s["header"], lead=s["subtitle"], sources=o["sources"], notes=s["notes"],
              footer="Public figures found by web search on 22 Oct 2026. Confirm IDC figures on idc.com.")
    return d


# =========================================================================== verify
def verify_pptx(path, d):
    """Re-open the .pptx and check slide count, notes, charts and hidden flags against the build."""
    prs = Presentation(path)
    out = {"slides": len(prs.slides), "expected": len(d.slides), "missing_notes": [], "charts": {},
           "hidden": [], "problems": []}
    for rec, sl in zip(d.slides, prs.slides):
        notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
        if not notes.strip():
            out["missing_notes"].append(rec.page)
        n_ch = sum(1 for sh in sl.shapes if sh.has_chart)
        if n_ch:
            out["charts"][rec.page] = n_ch
        if sl._element.get("show") == "0":
            out["hidden"].append(rec.page)
        want = sum(1 for e in rec.els if e.kind == "chart")
        if want != n_ch:
            out["problems"].append(f"slide {rec.page}: {n_ch} charts, expected {want}")
        if rec.hidden != (sl._element.get("show") == "0"):
            out["problems"].append(f"slide {rec.page}: hidden flag mismatch")
    if out["slides"] != out["expected"]:
        out["problems"].append("slide count mismatch")
    if out["missing_notes"]:
        out["problems"].append(f"slides without notes: {out['missing_notes']}")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--chrome", action="store_true", help="also run the headless Chrome overflow check")
    ap.add_argument("--shots", help="write a PNG per slide into this folder (headless Chrome)")
    a = ap.parse_args()
    d = build()
    d.save(OUT_PPTX)
    probs = d.check_layout()
    idx = d.html_preview(OUT_PREVIEW, problems=probs)
    if a.chrome:
        probs += chrome_check(idx)
    v = verify_pptx(OUT_PPTX, d)
    with open(OUT_REPORT, "w") as f:
        json.dump({"problems": probs, "pptx": v}, f, indent=1, ensure_ascii=False)
    if a.shots:
        os.makedirs(a.shots, exist_ok=True)
        for s in d.slides:
            chrome_screenshot(idx, s.index, os.path.join(a.shots, f"slide-{s.index:02d}.png"))
    errs = [p for p in probs if p.get("severity") == "error"]
    print(f"slides: {len(d.slides)} ({', '.join(s.page for s in d.slides)})\npptx: {OUT_PPTX}\npreview: {idx}")
    print(f"layout problems: {len(errs)} errors, {len(probs) - len(errs)} warnings")
    for p in probs:
        print(" ", p)
    print("pptx verify:", json.dumps(v, ensure_ascii=False))
    sys.exit(1 if errs or v["problems"] else 0)


if __name__ == "__main__":
    main()
